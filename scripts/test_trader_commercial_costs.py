"""Trader commercial costs, local AP and supplier ledger integration checks.

Use a local database with disposable test companies. --db enables SQL checks.
No FBR submission is performed. --keep leaves fixtures for browser verification.
"""
import argparse
import json
import io
import urllib.request
import uuid
import subprocess
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime, timezone
from pathlib import Path
from test_basic_flows import http, setup
from test_accounting_posting import sql


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--base", default="http://localhost:5288")
    ap.add_argument("--db", default=None)
    ap.add_argument("--keep", action="store_true")
    args = ap.parse_args()
    if args.db:
        parts = dict(part.strip().lower().split("=", 1) for part in args.db.split(";") if "=" in part)
        assert parts.get("server") == r".\mssqlserver02", "SQL checks require the local SQL instance"
        database = parts.get("database", "")
        assert database == "myapp_trader_local" or database.startswith("myapp_trader_costing_verify_"), "SQL checks require a local Trader database"
    token, company, client = setup(args.base, "admin", "admin123")
    cid = company["id"]
    today = datetime.now(timezone.utc).strftime("%Y-%m-%d")
    count = 0
    users = []
    other_company = None

    def call(method, path, body=None, expected=(200, 201), auth=token):
        status, data = http(method, path, args.base, token=auth, body=body)
        assert status in expected, f"{method} {path}: {status}: {data}"
        return data

    def check(name, condition):
        nonlocal count
        assert condition, name
        count += 1
        print(f"PASS {name}")

    def costs(kind, ident):
        return call("GET", f"/api/commercial-costs/{kind}/{ident}")

    def save(kind, ident, data):
        return call("PATCH", f"/api/commercial-costs/{kind}/{ident}", {
            "version": data["version"], "items": [{k: i.get(k) for k in ("id", "supplierId", "actualUnitCost")} for i in data["items"]]})

    try:
        call("POST", f"/api/accounts/company/{cid}/seed-wholesale")
        supplier = call("POST", "/api/suppliers", {"companyId": cid, "name": "Sample Cost Supplier A"})
        supplier_b = call("POST", "/api/suppliers", {"companyId": cid, "name": "Sample Cost Supplier B"})
        standalone = call("POST", "/api/invoices/standalone", {"companyId": cid, "clientId": client["id"], "date": today,
            "gstRate": 18, "items": [{"description": "Optional costing", "quantity": 2, "uom": "Pcs", "unitPrice": 100}]})
        bid = standalone["id"]
        preview = call("GET", f"/api/fbr/{bid}/preview-payload")
        preview_items = json.loads(preview["preview"]["json"])["items"] if preview.get("preview") else []
        check("unadjusted bill preview opens with missing HS code", preview["success"] and preview["preview"]["reviewRequired"] and preview_items and not preview_items[0]["hsCode"])
        for action in ("validate", "submit"):
            refused = call("POST", f"/api/fbr/{bid}/{action}")
            check(f"unadjusted bill {action} remains blocked before outbound posting", not refused["success"] and "Needs adjustment" in refused["errorMessage"])
        unchanged = call("GET", f"/api/invoices/{bid}")
        check("preview does not review or submit the commercial bill", unchanged["fbrAdjustmentRequired"] and not unchanged.get("fbrIRN"))
        d = costs("bill", bid)
        check("standalone creation without supplier or cost stays available", not d["costComplete"] and d["items"][0]["supplierId"] is None)
        d["items"][0]["actualUnitCost"] = 60
        d = save("bill", bid, d)
        check("actual cost does not require a supplier", d["actualAmount"] == 120 and d["profit"] == 80 and d["costComplete"])
        listed = call("GET", f"/api/invoices/company/{cid}/paged?pageSize=100")
        listed_bill = next(b for b in listed["items"] if b["id"] == bid)
        check("standalone list reports fully costed lines", listed_bill["costedLineCount"] == 1 and listed_bill["costingLineCount"] == 1)
        check("sales tax excluded from profit", d["sellingAmount"] == 200 and standalone["grandTotal"] == 236)
        d["items"][0]["supplierId"] = supplier["id"]
        body = {"version": d["version"], "items": [{k: i.get(k) for k in ("id", "supplierId", "actualUnitCost")} for i in d["items"]]}
        with ThreadPoolExecutor(max_workers=2) as pool:
            attempts = list(pool.map(lambda _: http("PATCH", f"/api/commercial-costs/bill/{bid}", args.base, token=token, body=body), range(2)))
        check("two simultaneous cost saves: one wins, the other is told to reload",
            sorted(status for status, _ in attempts) == [200, 400])
        auto = [p for p in call("GET", f"/api/purchasebills/company/{cid}/paged?pageSize=100")["items"] if p.get("sourceInvoiceId") == bid]
        check("saving a supplier and cost books exactly one supplier purchase", len(auto) == 1 and auto[0]["source"] == "auto-cost")
        purchase_id = auto[0]["id"]
        d = costs("bill", bid)
        save("bill", bid, d)
        check("saving the same costs again books nothing new",
            [p["id"] for p in call("GET", f"/api/purchasebills/company/{cid}/paged?pageSize=100")["items"] if p.get("sourceInvoiceId") == bid] == [purchase_id])
        p = call("GET", f"/api/purchasebills/{purchase_id}")
        check("local purchase is non-stock and retains bill provenance", p["purchaseCategory"] == "LocalDirect" and p["isNonStockSupplierPurchase"] and p["sourceInvoiceId"] == bid)
        ledger = call("GET", f"/api/commercial-costs/suppliers/{cid}/{supplier['id']}")
        check("supplier ledger has purchase liability", ledger["netBalance"] == 120)
        check("supplier purchase drills down to customer bill", ledger["purchases"][0]["invoiceIds"] == [bid])
        payment = call("POST", f"/api/payments/payments/company/{cid}", {"date": today, "contactType": "Supplier", "contactId": supplier["id"],
            "method": "Cash", "allocations": [{"purchaseBillId": purchase_id, "amount": 50}]})
        ledger = call("GET", f"/api/commercial-costs/suppliers/{cid}/{supplier['id']}")
        check("partial payment updates supplier balance", ledger["cashPaid"] == 50 and ledger["netBalance"] == 70)
        check("partial payment updates invoice-attributed outstanding", ledger["purchases"][0]["outstanding"] == 70)
        panel = call("GET", f"/api/commercial-costs/bill/{bid}/payables")
        check("bill payable panel reconciles to supplier ledger", panel[0]["outstanding"] == 70)
        check("payment does not change commercial profit", costs("bill", bid)["profit"] == 80)
        if args.db:
            # Hold a different document's allocation uncommitted. A clustered
            # scan in receipt recalculation used to wait on this supplier row.
            lock_query = f"""SET NOCOUNT ON; BEGIN TRAN;
                UPDATE PaymentAllocations WITH (ROWLOCK) SET Amount=Amount WHERE PaymentId={payment['id']};
                RAISERROR('LOCK_READY',0,1) WITH NOWAIT;
                WAITFOR DELAY '00:00:06'; ROLLBACK;"""
            lock = subprocess.Popen(["sqlcmd", "-S", parts["server"], "-d", database,
                "-E", "-C", "-t", "15", "-Q", lock_query], stdout=subprocess.PIPE,
                stderr=subprocess.STDOUT, text=True)
            try:
                while True:
                    line = lock.stdout.readline()
                    assert line, "allocation lock session ended before readiness"
                    if "LOCK_READY" in line:
                        break
                lock_request = urllib.request.Request(args.base + f"/api/payments/receipts/company/{cid}",
                    data=json.dumps({"date": today, "contactType": "Client", "contactId": client["id"],
                        "method": "Cash", "allocations": [{"invoiceId": bid, "amount": 1}]}).encode(),
                    headers={"Content-Type": "application/json", "Authorization": "Bearer " + token}, method="POST")
                with urllib.request.urlopen(lock_request, timeout=3) as response:
                    check("receipt completes while unrelated supplier allocation is locked", response.status == 201)
            finally:
                lock.communicate(timeout=20)

        call("POST", f"/api/payments/payments/company/{cid}", {"date": today, "contactType": "Supplier", "contactId": supplier["id"],
            "method": "Cash", "allocations": [{"kind": "OnAccount", "amount": 10}]})
        ledger = call("GET", f"/api/commercial-costs/suppliers/{cid}/{supplier['id']}")
        check("advance is visible and affects net supplier balance once", ledger["unappliedAdvances"] == 10 and ledger["netBalance"] == 60)
        check("unapplied advance does not pretend to settle a purchase", ledger["purchases"][0]["outstanding"] == 70)
        d = costs("bill", bid)
        old_version = d["version"]
        d["items"][0]["actualUnitCost"] = 70
        d = save("bill", bid, d)
        check("estimate can change after settlement without re-pricing the recorded purchase", d["profit"] == 80 and d["items"][0]["purchasedAmount"] == 120)
        ledger = call("GET", f"/api/commercial-costs/suppliers/{cid}/{supplier['id']}")
        check("cost correction preserves settled liability and flags review", ledger["netBalance"] == 60 and ledger["purchases"][0]["costChanged"])
        call("PATCH", f"/api/commercial-costs/bill/{bid}", {"version": old_version, "items": d["items"]}, expected=(400,))
        check("stale cost edit cannot overwrite a newer correction", True)
        printed = call("GET", f"/api/invoices/{bid}/print/bill")
        check("customer print does not expose private cost", "actualunitcost" not in json.dumps(printed).lower())
        if args.db:
            sql(args.db, f"UPDATE Invoices SET FbrStatus='Submitted', FbrSubmittedAt=GETUTCDATE(), FbrIRN='SAMPLE-COST-TEST' WHERE Id={bid} AND CompanyId={cid}")
            before = call("GET", f"/api/invoices/{bid}")
            before_print = call("GET", f"/api/invoices/{bid}/print/bill")
            d = costs("bill", bid); d["items"][0]["actualUnitCost"] = 75
            d = save("bill", bid, d)
            after = call("GET", f"/api/invoices/{bid}")
            check("filed standalone actual cost remains editable", d["items"][0]["actualUnitCost"] == 75 and d["profit"] == 80)
            check("filed sale values and FBR identity remain unchanged", all(before[k] == after[k] for k in ("subtotal", "gstAmount", "grandTotal", "fbrIRN", "fbrStatus", "items")))
            check("filed customer print remains unchanged", before_print == call("GET", f"/api/invoices/{bid}/print/bill"))
            audit = sql(args.db, f"SET NOCOUNT ON; SELECT COUNT(*) FROM PrivateCostChanges WHERE CompanyId={cid} AND Target='bill' AND TargetId={bid}")
            check("private cost corrections are audited", int(audit.strip()) >= 4)
            stock = sql(args.db, f"SET NOCOUNT ON; SELECT COUNT(*) FROM StockMovements WHERE CompanyId={cid} AND SourceType=1 AND SourceId={purchase_id}")
            check("local purchase creates no inventory movement", int(stock.strip()) == 0)

        c = call("POST", f"/api/deliverychallans/company/{cid}", {"clientId": client["id"], "poNumber": "SAMPLE-COST-PO", "deliveryDate": today,
            "items": [{"description": "Challan A", "quantity": 3, "unit": "Pcs", "actualUnitCost": 20, "supplierId": supplier["id"]},
                {"description": "Challan B", "quantity": 2, "unit": "Pcs"}]})
        cd = costs("challan", c["id"])
        check("a costed challan line is booked on save while a blank line stays open",
            cd["items"][0]["remainingQuantity"] == 0 and cd["items"][1]["remainingQuantity"] == 2)
        cd["items"][1]["supplierId"] = supplier_b["id"]; cd["items"][1]["actualUnitCost"] = 30
        cd = save("challan", c["id"], cd)
        check("a line costed later is booked by that save", cd["items"][1]["remainingQuantity"] == 0)

        split = call("POST", f"/api/deliverychallans/company/{cid}", {"clientId": client["id"], "poNumber": "SAMPLE-SPLIT-PURCHASE", "deliveryDate": today,
            "items": [{"description": "Split A1", "quantity": 4, "unit": "Pcs"}, {"description": "Split A2", "quantity": 3, "unit": "Pcs"}, {"description": "Split B1", "quantity": 2, "unit": "Pcs"}]})
        split_items = split["items"]
        picker = call("GET", f"/api/purchasebills/company/{cid}/local-purchase-sources")
        check("local purchase picker is company-scoped and lists unpurchased challans", any(x["kind"] == "challan" and x["id"] == split["id"] and x["remainingLines"] == 3 for x in picker))
        selected = call("GET", f"/api/purchasebills/company/{cid}/local-purchase-sources/challan/{split['id']}")
        check("challan picker returns each available line and quantity", len(selected["items"]) == 3 and all(x["remainingQuantity"] == x["quantity"] for x in selected["items"]))
        def local_bill(supplier_id, allocations, label):
            return call("POST", "/api/purchasebills", {"companyId": cid, "supplierId": supplier_id, "date": today,
                "purchaseCategory": "LocalDirect", "gstRate": 0, "supplierBillNumber": label,
                "items": [{"description": split_items[index]["description"], "quantity": qty, "uom": "Pcs", "unitPrice": price,
                    "localDeliveryItemId": split_items[index]["id"]} for index, qty, price in allocations]})
        split_a_first = local_bill(supplier["id"], [(0, 2, 11), (1, 2, 21)], "SPLIT-A-1")
        partial_costs = costs("challan", split["id"])
        check("a line-linked bill leaves the challan's other lines open in costing",
            [x["remainingQuantity"] for x in partial_costs["items"]] == [2, 1, 2]
            and not partial_costs["items"][2]["purchaseRecorded"] and partial_costs["items"][2]["purchasedAmount"] == 0)
        split_b = local_bill(supplier_b["id"], [(2, 2, 31)], "SPLIT-B")
        check("one supplier can purchase multiple selected challan items", len(split_a_first["items"]) == 2 and split_a_first["sourceDeliveryChallanId"] == split["id"])
        check("another supplier can purchase a different item from the same challan", len(split_b["items"]) == 1 and split_b["supplierId"] == supplier_b["id"])
        check("partial purchase allocations reduce available quantities", call("GET", f"/api/purchasebills/company/{cid}/local-purchase-sources/challan/{split['id']}")["items"][0]["remainingQuantity"] == 2)
        call("POST", "/api/purchasebills", {"companyId": cid, "supplierId": supplier["id"], "date": today, "purchaseCategory": "LocalDirect", "gstRate": 0,
            "items": [{"description": "Over allocation", "quantity": 3, "unitPrice": 12, "localDeliveryItemId": split_items[0]["id"]}]}, expected=(400,))
        check("over-allocation is rejected by the server", True)
        split_a_second = local_bill(supplier["id"], [(0, 2, 12), (1, 1, 22)], "SPLIT-A-2")
        check("a source line can span multiple purchase bills and suppliers", len(split_a_second["items"]) == 2)
        split_costs = costs("challan", split["id"])
        check("challan costing aggregates split purchase amounts and remaining quantities",
            [x["purchasedAmount"] for x in split_costs["items"]] == [46, 64, 62]
            and [x["remainingQuantity"] for x in split_costs["items"]] == [0, 0, 0]
            and split_costs["actualAmount"] == 172 and split_costs["costComplete"])
        split_invoice = call("POST", "/api/invoices", {"companyId": cid, "clientId": client["id"], "date": today, "gstRate": 0,
            "challanIds": [split["id"]], "items": [{"deliveryItemId": i["id"], "unitPrice": 50} for i in split_items]})
        check("bill profit includes purchases created across suppliers", costs("bill", split_invoice["id"])["profit"] == 278)
        split_payables = call("GET", f"/api/commercial-costs/bill/{split_invoice['id']}/payables")
        check("bill payable drill-down includes each split supplier purchase", len(split_payables) == 3 and sum(p["liability"] for p in split_payables) == 172)


        # ── One pool, two entry points: challan or sale bill, form or dialog ──
        pool_bill = call("POST", "/api/invoices/standalone", {"companyId": cid, "clientId": client["id"], "date": today, "gstRate": 0,
            "items": [{"description": "Pool fitting", "quantity": 5, "uom": "Pcs", "unitPrice": 30},
                      {"description": "Pool hose", "quantity": 2, "uom": "Pcs", "unitPrice": 80}]})
        pool_items = pool_bill["items"]
        sources = call("GET", f"/api/purchasebills/company/{cid}/local-purchase-sources")
        check("picker lists sale bills with quantity left and drops fully bought documents",
            any(x["kind"] == "bill" and x["id"] == pool_bill["id"] and x["remainingLines"] == 2 for x in sources)
            and not any(x["kind"] == "challan" and x["id"] == split["id"] for x in sources))
        found = call("GET", f"/api/purchasebills/company/{cid}/local-purchase-sources?search={pool_bill['invoiceNumber']}")
        check("picker search finds a sale bill by number", any(x["kind"] == "bill" and x["id"] == pool_bill["id"] for x in found))
        pool_source = call("GET", f"/api/purchasebills/company/{cid}/local-purchase-sources/bill/{pool_bill['id']}")
        check("standalone bill lines are offered as bill lines",
            [x["invoiceItemId"] for x in pool_source["items"]] == [i["id"] for i in pool_items] and all(x["deliveryItemId"] is None for x in pool_source["items"]))

        def local_purchase(supplier_id, lines, expected=(200, 201)):
            return call("POST", "/api/purchasebills", {"companyId": cid, "supplierId": supplier_id, "date": today,
                "purchaseCategory": "LocalDirect", "gstRate": 0, "items": [dict({"description": "Pool line", "uom": "Pcs"}, **line) for line in lines]}, expected=expected)

        pool_first = local_purchase(supplier["id"], [{"quantity": 3, "unitPrice": 12, "localInvoiceItemId": pool_items[0]["id"]}])
        check("form links a purchase to a standalone bill line", pool_first["items"][0]["localInvoiceItemId"] == pool_items[0]["id"] and pool_first["sourceInvoiceId"] == pool_bill["id"])
        pd = costs("bill", pool_bill["id"])
        check("bill costing sees the form purchase", pd["items"][0]["purchasedQuantity"] == 3 and pd["items"][0]["remainingQuantity"] == 2 and pd["items"][0]["purchasedAmount"] == 36)
        hose = pd["items"][1]["id"]
        local_purchase(supplier_b["id"], [{"quantity": 1, "unitPrice": 50, "localInvoiceItemId": hose}])
        pd = costs("bill", pool_bill["id"])
        check("a partial supplier bill records only that quantity", pd["items"][1]["purchasedQuantity"] == 1 and pd["items"][1]["remainingQuantity"] == 1 and pd["items"][1]["purchasedAmount"] == 50)
        local_purchase(supplier["id"], [{"quantity": 2, "unitPrice": 45, "localInvoiceItemId": hose}], expected=(400,))
        check("a supplier bill cannot buy more than is left", True)
        local_purchase(supplier["id"], [{"quantity": 1, "unitPrice": 1, "localInvoiceItemId": pool_items[0]["id"]},
                                        {"quantity": 1, "unitPrice": 1, "localInvoiceItemId": pool_items[0]["id"]}], expected=(400,))
        check("one purchase cannot claim a source line twice", True)
        call("POST", "/api/purchasebills", {"companyId": cid, "supplierId": supplier["id"], "date": today, "purchaseCategory": "Inventory", "gstRate": 0,
            "items": [{"description": "Inventory link", "quantity": 1, "unitPrice": 1, "localInvoiceItemId": hose}]}, expected=(400,))
        check("inventory purchases cannot claim a sale line", True)

        # A race: two saves for the last unit, exactly one wins.
        with ThreadPoolExecutor(max_workers=2) as pool:
            raced = list(pool.map(lambda _: http("POST", "/api/purchasebills", args.base, token=token, body={"companyId": cid, "supplierId": supplier["id"],
                "date": today, "purchaseCategory": "LocalDirect", "gstRate": 0,
                "items": [{"description": "Race", "quantity": 1, "unitPrice": 5, "uom": "Pcs", "localInvoiceItemId": hose}]})[0], range(2)))
        check("two simultaneous saves cannot both spend the last unit", sorted(raced) == [201, 400] or sorted(raced) == [200, 400])
        check("pool is exhausted exactly once", costs("bill", pool_bill["id"])["items"][1]["remainingQuantity"] == 0)

        # Correcting an unpaid linked purchase, then the lock once it is paid.
        def edit_payload(bill, quantity):
            return {"supplierId": bill["supplierId"], "date": today, "gstRate": 0, "items": [{"id": i["id"], "description": i["description"],
                "quantity": quantity, "uom": i["uom"], "unitPrice": i["unitPrice"]} for i in bill["items"]]}
        edited = call("PUT", f"/api/purchasebills/{pool_first['id']}", edit_payload(pool_first, 5))
        check("an unpaid linked purchase can be corrected within what is left",
            edited["items"][0]["quantity"] == 5 and edited["items"][0]["localInvoiceItemId"] == pool_items[0]["id"] and edited["grandTotal"] == 60)
        call("PUT", f"/api/purchasebills/{pool_first['id']}", edit_payload(edited, 6), expected=(400,))
        check("a correction cannot claim more than the sale line holds", True)
        edited_again = call("PUT", f"/api/purchasebills/{pool_first['id']}", edit_payload(edited, 4))
        check("lowering a purchase hands quantity back to the pool", costs("bill", pool_bill["id"])["items"][0]["remainingQuantity"] == 1)

        # The sale side cannot drop below what suppliers were paid for.
        pool_invoice = call("GET", f"/api/invoices/{pool_bill['id']}")
        lowered = {"gstRate": 0, "items": [{"id": i["id"], "description": i["description"], "quantity": 3 if i["id"] == pool_items[0]["id"] else float(i["quantity"]),
            "uom": i.get("uom") or "Pcs", "unitPrice": float(i["unitPrice"])} for i in pool_invoice["items"]]}
        call("PUT", f"/api/invoices/{pool_bill['id']}", lowered, expected=(400,))
        check("a sale line cannot be lowered below the quantity already purchased", call("GET", f"/api/invoices/{pool_bill['id']}")["items"][0]["quantity"] == 5)
        lowered["items"][0]["quantity"] = 4
        call("PUT", f"/api/invoices/{pool_bill['id']}", lowered)
        check("a sale line can be lowered down to the purchased quantity", call("GET", f"/api/invoices/{pool_bill['id']}")["items"][0]["quantity"] == 4)
        call("DELETE", f"/api/invoices/{pool_bill['id']}", expected=(400,))
        check("a bill with supplier purchases cannot be deleted", True)

        call("POST", f"/api/payments/payments/company/{cid}", {"date": today, "contactType": "Supplier", "contactId": supplier["id"],
            "method": "Cash", "allocations": [{"purchaseBillId": pool_first["id"], "amount": 10}]})
        call("PUT", f"/api/purchasebills/{pool_first['id']}", edit_payload(edited_again, 3), expected=(400,))
        check("a paid linked purchase is corrected only by a supplier credit", call("GET", f"/api/purchasebills/{pool_first['id']}")["items"][0]["quantity"] == 4)

        # A bill made from a challan offers, and books against, the CHALLAN line.
        via = call("POST", f"/api/deliverychallans/company/{cid}", {"clientId": client["id"], "poNumber": "SAMPLE-VIA-BILL", "deliveryDate": today,
            "items": [{"description": "Via bill bracket", "quantity": 6, "unit": "Pcs"}]})
        via_bill = call("POST", "/api/invoices", {"companyId": cid, "clientId": client["id"], "date": today, "gstRate": 0,
            "challanIds": [via["id"]], "items": [{"deliveryItemId": via["items"][0]["id"], "unitPrice": 20}]})
        via_source = call("GET", f"/api/purchasebills/company/{cid}/local-purchase-sources/bill/{via_bill['id']}")
        check("a challan-backed bill line is offered as its challan line",
            via_source["items"][0]["deliveryItemId"] == via["items"][0]["id"] and via_source["items"][0]["invoiceItemId"] is None and via_source["items"][0]["sourceChallanNumber"] == via["challanNumber"])
        via_buy = local_purchase(supplier["id"], [{"quantity": 2, "unitPrice": 9, "localInvoiceItemId": via_bill["items"][0]["id"]}])
        check("a purchase sent against the bill line is booked to the challan line",
            via_buy["items"][0]["localDeliveryItemId"] == via["items"][0]["id"] and via_buy["items"][0]["localInvoiceItemId"] is None)
        check("challan and bill views draw on the same remaining quantity",
            costs("challan", via["id"])["items"][0]["remainingQuantity"] == 4 and costs("bill", via_bill["id"])["items"][0]["remainingQuantity"] == 4)
        local_purchase(supplier_b["id"], [{"quantity": 5, "unitPrice": 9, "localDeliveryItemId": via["items"][0]["id"]}], expected=(400,))
        check("the challan view cannot re-buy what was bought against the bill", True)
        bill_payables = call("GET", f"/api/commercial-costs/bill/{via_bill['id']}/payables")
        check("a purchase bought against the bill shows in the bill's payables", any(p["id"] == via_buy["id"] for p in bill_payables))
        check("a partly bought challan still lists in the picker",
            any(x["kind"] == "challan" and x["id"] == via["id"] and x["remainingQuantity"] == 4 for x in call("GET", f"/api/purchasebills/company/{cid}/local-purchase-sources")))
        call("GET", f"/api/purchasebills/company/{cid}/local-purchase-sources/invoice/{via_bill['id']}", expected=(404,))
        check("unknown source kind is refused", True)

        linked = call("POST", "/api/invoices", {"companyId": cid, "clientId": client["id"], "date": today, "gstRate": 0,
            "challanIds": [c["id"]], "items": [{"deliveryItemId": i["id"], "unitPrice": 50} for i in c["items"]]})
        ld = costs("bill", linked["id"])
        check("challan bill uses source actual costs", ld["actualAmount"] == 120 and ld["profit"] == 130)
        listed = call("GET", f"/api/invoices/company/{cid}/paged?pageSize=100")
        listed_bill = next(b for b in listed["items"] if b["id"] == linked["id"])
        check("linked bill list counts current challan costs", listed_bill["costedLineCount"] == 2 and listed_bill["costingLineCount"] == 2)
        linked_payables = call("GET", f"/api/commercial-costs/bill/{linked['id']}/payables")
        check("purchases created before billing automatically appear on bill", len(linked_payables) == 2 and sum(p["liability"] for p in linked_payables) == 120)
        if args.db:
            sql(args.db, f"UPDATE Invoices SET FbrStatus='Submitted', FbrSubmittedAt=GETUTCDATE(), FbrIRN='SAMPLE-CHALLAN-TEST' WHERE Id={linked['id']} AND CompanyId={cid}")
        cd = costs("challan", c["id"]); cd["items"][0]["actualUnitCost"] = 25
        save("challan", c["id"], cd)
        check("billed/filed challan cost remains editable", costs("challan", c["id"])["items"][0]["actualUnitCost"] == 25 and costs("bill", linked["id"])["profit"] == 115)
        check("an unpaid booked purchase follows the corrected cost", sum(p["liability"] for p in call("GET", f"/api/commercial-costs/bill/{linked['id']}/payables")) == 135)

        direct = call("POST", "/api/invoices/standalone", {"companyId": cid, "clientId": client["id"], "date": today, "gstRate": 0,
            "items": [{"description": "Cost at creation", "quantity": 2, "uom": "Pcs", "unitPrice": 100, "actualUnitCost": 40, "supplierId": supplier["id"]}]})
        check("standalone creation retains optional actual cost", costs("bill", direct["id"])["profit"] == 120)
        d = costs("bill", direct["id"]); d["items"][0]["actualUnitCost"] = 0; d["items"][0]["supplierId"] = None
        check("explicit zero cost is complete without supplier", save("bill", direct["id"], d)["costComplete"])
        d = costs("bill", direct["id"]); d["items"][0]["actualUnitCost"] = None
        check("cost and supplier can be cleared without blocking bill", not save("bill", direct["id"], d)["costComplete"])
        d = costs("bill", direct["id"]); d["items"][0]["actualUnitCost"] = -1
        call("PATCH", f"/api/commercial-costs/bill/{direct['id']}", {"version": d["version"], "items": d["items"]}, expected=(400,))
        check("negative costs refused", True)

        combined_challans = []
        for qty, cost in [(2, 10), (3, 20)]:
            combined_challans.append(call("POST", f"/api/deliverychallans/company/{cid}", {"clientId": client["id"], "poNumber": "SAMPLE-COMBINED", "deliveryDate": today,
                "items": [{"description": "Combined cost sample", "quantity": qty, "unit": "Pcs", "actualUnitCost": cost, "supplierId": supplier_b["id"]}]}))
        combined = call("POST", "/api/invoices", {"companyId": cid, "clientId": client["id"], "date": today, "gstRate": 0,
            "challanIds": [c["id"] for c in combined_challans], "items": [{"deliveryItemId": c["items"][0]["id"], "unitPrice": 50} for c in combined_challans]})
        cd = costs("bill", combined["id"])
        check("bill sums actual costs across multiple challans", cd["actualAmount"] == 80 and cd["profit"] == 170)
        combined_purchases = [p["id"] for p in call("GET", f"/api/purchasebills/company/{cid}/paged?pageSize=100")["items"]
            if p.get("sourceDeliveryChallanId") in [c["id"] for c in combined_challans]]
        check("combined bill keeps one purchase per source challan and supplier", len(combined_purchases) == 2)
        for ident in combined_purchases:
            pb = call("GET", f"/api/purchasebills/{ident}")
            call("PUT", f"/api/purchasebills/{ident}", {"supplierId": pb["supplierId"], "date": pb["date"], "gstRate": 10,
                "supplierBillNumber": f"SAMPLE-{cid}-{ident}", "items": [{"id": i["id"], "description": i["description"], "quantity": i["quantity"], "uom": i["uom"], "unitPrice": i["unitPrice"]} for i in pb["items"]]})
        purchases_with_tax = [call("GET", f"/api/purchasebills/{ident}") for ident in combined_purchases]
        check("supplier preview tax and references flow to purchase bills", sum(p["grandTotal"] for p in purchases_with_tax) == 88 and all(p["supplierBillNumber"] for p in purchases_with_tax))
        check("supplier purchase tax does not inflate commercial actual cost", costs("bill", combined["id"])["profit"] == 170)

        manual = call("POST", "/api/purchasebills", {"companyId": cid, "supplierId": supplier_b["id"], "date": today, "purchaseCategory": "LocalDirect", "gstRate": 0,
            "supplierBillNumber": "COST-REF-SAMPLE", "items": [{"description": "Local import reconciliation", "quantity": 2, "uom": "Pcs", "unitPrice": 10}]})
        check("manual local purchase needs no inventory classification", manual["purchaseCategory"] == "LocalDirect" and manual["grandTotal"] == 20)
        if args.db:
            sql(args.db, f"UPDATE Companies SET InventoryTrackingEnabled=1 WHERE Id={cid}")
            tracked_local = call("POST", "/api/purchasebills", {"companyId": cid, "supplierId": supplier_b["id"], "date": today, "purchaseCategory": "LocalDirect", "gstRate": 0,
                "items": [{"description": "Local while stock tracking enabled", "quantity": 1, "uom": "Pcs", "unitPrice": 25}]})
            stock = sql(args.db, f"SET NOCOUNT ON; SELECT COUNT(*) FROM StockMovements WHERE CompanyId={cid} AND SourceType=1 AND SourceId={tracked_local['id']}")
            check("local purchase does not move stock in inventory-enabled company", int(stock.strip()) == 0)
            gl = sql(args.db, f"SET NOCOUNT ON; SELECT COUNT(*) FROM JournalLines l JOIN JournalEntries e ON e.Id=l.JournalEntryId JOIN Accounts a ON a.Id=l.AccountId WHERE e.CompanyId={cid} AND e.SourceDocType=2 AND e.SourceDocId={tracked_local['id']} AND a.Name='Cost of goods sold' AND l.Debit=25")
            check("local purchase posts to cost rather than inventory", int(gl.strip()) == 1)

        from openpyxl import Workbook
        workbook = Workbook(); sheet = workbook.active; sheet.title = "Domestic Invoices"
        sheet.append(["Invoice Ref No.", "Status", "Invoice No.", "Invoice Type", "Invoice Date", "Seller Registration No.", "Seller Name", "Taxpayer Type", "Sale Type", "Quantity", "Product Description", "HS Code", "Rate", "UoM", "Value of Sales Excluding Sales Tax", "Sales Tax/ FED in ST Mode", "Total Value of Sales"])
        sheet.append(["COST-REF-SAMPLE", "Valid", "COST-REF-SAMPLE", "Purchase Invoice", today, "1112223", supplier_b["name"], "Registered", "Goods at standard rate (default)", 2, "Local import reconciliation", "7306.9000", "0%", "PCS", 20, 0, 20])
        stream = io.BytesIO(); workbook.save(stream)

        def upload(path):
            boundary = "costing" + uuid.uuid4().hex
            body = (f"--{boundary}\r\nContent-Disposition: form-data; name=\"companyId\"\r\n\r\n{cid}\r\n--{boundary}\r\nContent-Disposition: form-data; name=\"file\"; filename=\"sample.xlsx\"\r\nContent-Type: application/vnd.openxmlformats-officedocument.spreadsheetml.sheet\r\n\r\n".encode()
                + stream.getvalue() + f"\r\n--{boundary}--\r\n".encode())
            request = urllib.request.Request(args.base + path, data=body, headers={"Authorization": "Bearer " + token, "Content-Type": "multipart/form-data; boundary=" + boundary})
            with urllib.request.urlopen(request) as response: return json.load(response)

        preview = upload("/api/fbr-purchase-import/preview")
        check("FBR import identifies existing local purchase", preview["invoices"][0]["matchedPurchaseBillId"] == manual["id"])
        upload("/api/fbr-purchase-import/commit")
        reconciled = call("GET", f"/api/purchasebills/{manual['id']}")
        check("FBR match attaches tax evidence without changing local category", reconciled["reconciliationStatus"] == "Matched" and reconciled["purchaseCategory"] == "LocalDirect")
        all_purchases = call("GET", f"/api/purchasebills/company/{cid}/paged?pageSize=100")
        check("FBR reconciliation creates no duplicate payable", sum(1 for p in all_purchases["items"] if p.get("supplierBillNumber") == "COST-REF-SAMPLE") == 1)
        local_view = call("GET", f"/api/purchasebills/company/{cid}/paged?purchaseCategory=LocalDirect&pageSize=100")
        check("local purchase view excludes inventory purchases", all(p["purchaseCategory"] == "LocalDirect" for p in local_view["items"]))
        sheet.cell(2, 1, "SAMPLE-NEW-FBR"); sheet.cell(2, 3, "SAMPLE-NEW-FBR")
        stream = io.BytesIO(); workbook.save(stream)
        imported = upload("/api/fbr-purchase-import/commit")
        check("unmatched FBR purchase imports once", imported["counts"]["invoicesImported"] == 1)
        fbr_view = call("GET", f"/api/purchasebills/company/{cid}/paged?entrySource=fbr-import&pageSize=100")
        check("FBR purchase view stays separate from local purchases", len(fbr_view["items"]) == 1 and fbr_view["items"][0]["purchaseCategory"] == "Inventory")
        if args.db:
            imported_id = fbr_view["items"][0]["id"]
            stock = sql(args.db, f"SET NOCOUNT ON; SELECT COUNT(*) FROM StockMovements WHERE CompanyId={cid} AND SourceType=1 AND SourceId={imported_id}")
            check("inventory FBR import records its stock receipt", int(stock.strip()) == 1)
            entry = sql(args.db, f"SET NOCOUNT ON; SELECT COUNT(*) FROM JournalEntries WHERE CompanyId={cid} AND SourceDocType=2 AND SourceDocId={imported_id}")
            check("FBR purchase import posts its supplier liability to GL", int(entry.strip()) == 1)
        before_credit = call("GET", f"/api/commercial-costs/suppliers/{cid}/{supplier_b['id']}")
        call("POST", "/api/purchasedebitnotes", {"companyId": cid, "supplierId": supplier_b["id"], "date": today, "gstRate": 0,
            "items": [{"description": "Sample supplier credit", "quantity": 1, "uom": "Pcs", "unitPrice": 5}]})
        after_credit = call("GET", f"/api/commercial-costs/suppliers/{cid}/{supplier_b['id']}")
        check("supplier credit reduces net ledger balance once", after_credit["netBalance"] == before_credit["netBalance"] - 5 and after_credit["unallocatedCredits"] == 5)
        linked_credit = call("POST", "/api/purchasedebitnotes", {"companyId": cid, "supplierId": supplier["id"], "purchaseBillId": purchase_id, "date": today, "gstRate": 0,
            "items": [{"description": "Credit linked to original supplier purchase", "quantity": 1, "uom": "Pcs", "unitPrice": 5}]})
        linked_balance = next(p for p in call("GET", f"/api/commercial-costs/suppliers/{cid}/{supplier['id']}")["purchases"] if p["id"] == purchase_id)
        check("linked supplier credit reduces the correct invoice outstanding", linked_balance["creditAmount"] == 5 and linked_balance["outstanding"] == 65)
        check("payment picker respects supplier credits", call("GET", f"/api/purchasebills/{purchase_id}")["balanceDue"] == 65)
        call("POST", f"/api/payments/payments/company/{cid}", {"date": today, "contactType": "Supplier", "contactId": supplier["id"], "method": "Cash",
            "allocations": [{"purchaseBillId": purchase_id, "amount": 66}]}, expected=(400,))
        check("payment cannot overpay a credited purchase", True)
        call("DELETE", f"/api/purchasedebitnotes/{linked_credit['id']}", expected=(200, 204))
        check("reversing supplier credit restores purchase outstanding", call("GET", f"/api/purchasebills/{purchase_id}")["balanceDue"] == 70)
        historical = call("GET", f"/api/commercial-costs/suppliers/{cid}/{supplier_b['id']}?asOf=2000-01-01")
        check("as-of ledger excludes later purchases and settlements", historical["netBalance"] == 0 and not historical["rows"])
        if args.db:
            sql(args.db, f"UPDATE Companies SET GlLockDate='{today}' WHERE Id={cid}")
            d = costs("bill", bid); d["items"][0]["actualUnitCost"] = 76
            check("closed accounting period does not block private costing correction", save("bill", bid, d)["items"][0]["actualUnitCost"] == 76)

        other_company = call("POST", "/api/companies", {"name": f"Sample other cost company {cid}"})
        foreign = call("POST", "/api/suppliers", {"companyId": other_company["id"], "name": "Sample other supplier"})
        call("POST", "/api/purchasedebitnotes", {"companyId": other_company["id"], "supplierId": foreign["id"], "purchaseBillId": purchase_id, "date": today, "gstRate": 0,
            "items": [{"description": "Forbidden cross-company purchase credit", "quantity": 1, "uom": "Pcs", "unitPrice": 5}]}, expected=(400,))
        check("supplier credit cannot reference another company's purchase", True)
        d = costs("bill", direct["id"]); d["items"][0]["supplierId"] = foreign["id"]
        call("PATCH", f"/api/commercial-costs/bill/{direct['id']}", {"version": d["version"], "items": d["items"]}, expected=(400,))
        check("cross-company supplier rejected on cost update", True)
        call("POST", "/api/purchasebills", {"companyId": other_company["id"], "supplierId": foreign["id"], "date": today,
            "purchaseCategory": "LocalDirect", "gstRate": 0, "items": [{"description": "Foreign challan link", "quantity": 1,
                "unitPrice": 5, "localDeliveryItemId": c["items"][0]["id"]}]}, expected=(400,))
        check("cross-company challan line cannot be linked to a local purchase", True)
        call("POST", "/api/purchasebills", {"companyId": other_company["id"], "supplierId": foreign["id"], "date": today,
            "purchaseCategory": "LocalDirect", "gstRate": 0, "items": [{"description": "Foreign bill link", "quantity": 1,
                "unitPrice": 5, "localInvoiceItemId": direct["items"][0]["id"]}]}, expected=(400,))
        check("cross-company bill line cannot be linked to a local purchase", True)
        for kind, ident in (("challan", c["id"]), ("bill", bid)):
            call("GET", f"/api/purchasebills/company/{other_company['id']}/local-purchase-sources/{kind}/{ident}", expected=(404,))
            check(f"a {kind} is not reachable through another company's picker route", True)
        check("another company's picker lists none of this company's documents",
            not any(x["id"] in (c["id"], bid) for x in call("GET", f"/api/purchasebills/company/{other_company['id']}/local-purchase-sources")))
        roles = call("GET", "/api/roles")
        sales = next(r["id"] for r in roles if r["name"] == "Sales Edition")
        username = f"costverify{cid}"
        u = call("POST", "/api/users", {"username": username, "password": "Testcost1234", "fullName": "Sample cost tester", "role": "User"}); users.append(u)
        call("PUT", f"/api/users/{u['id']}/roles", {"roleIds": [sales]})
        call("PUT", f"/api/usercompanies/user/{u['id']}", {"companyIds": [other_company["id"]]})
        user_token = call("POST", "/api/auth/login", {"username": username, "password": "Testcost1234"})["token"]
        for method, path, body in [
            ("GET", f"/api/fbr/{bid}/preview-payload", None),
            ("GET", f"/api/invoices/company/{cid}/paged?costingStatus=full", None),
            ("GET", f"/api/deliverychallans/company/{cid}/paged?costingStatus=partial", None),
            ("GET", f"/api/purchasebills/company/{cid}/local-purchase-sources", None),
            ("GET", f"/api/purchasebills/company/{cid}/local-purchase-sources/challan/{c['id']}", None),
            ("GET", f"/api/purchasebills/company/{cid}/local-purchase-sources/bill/{bid}", None),
            ("GET", f"/api/commercial-costs/bill/{bid}", None),
            ("PATCH", f"/api/commercial-costs/bill/{bid}", {"version": "forged", "items": []}),
            ("GET", f"/api/commercial-costs/challan/{c['id']}", None),
            ("PATCH", f"/api/commercial-costs/challan/{c['id']}", {"version": "forged", "items": []}),
            ("GET", f"/api/commercial-costs/bill/{bid}/payables", None),
            ("GET", f"/api/commercial-costs/suppliers/{cid}", None),
            ("GET", f"/api/commercial-costs/suppliers/{cid}/{supplier['id']}", None),
        ]:
            call(method, path, body, expected=(403,), auth=user_token)
            check(f"cross-company refusal: {method} {path.split('/api/')[1]}", True)
        call("GET", f"/api/commercial-costs/suppliers/{other_company['id']}", auth=user_token)
        check("Sales Edition can read its own supplier ledger summary", True)
        call("PUT", f"/api/usercompanies/user/{u['id']}", {"companyIds": [cid]})
        user_token = call("POST", "/api/auth/login", {"username": username, "password": "Testcost1234"})["token"]
        check("Sales Edition can read its assigned commercial cost document", call("GET", f"/api/commercial-costs/bill/{bid}", auth=user_token)["id"] == bid)
        call("PUT", f"/api/users/{u['id']}/roles", {"roleIds": []})
        user_token = call("POST", "/api/auth/login", {"username": username, "password": "Testcost1234"})["token"]
        call("GET", f"/api/commercial-costs/bill/{bid}", auth=user_token, expected=(403,))
        call("PATCH", f"/api/commercial-costs/bill/{bid}", {"version": "forged", "items": []}, auth=user_token, expected=(403,))
        call("GET", f"/api/commercial-costs/suppliers/{cid}", auth=user_token, expected=(403,))
        check("company assignment alone grants neither cost nor payable permissions", True)
        summary = call("GET", f"/api/commercial-costs/suppliers/{cid}")
        a = next(s for s in summary if s["supplierId"] == supplier["id"])
        ledger = call("GET", f"/api/commercial-costs/suppliers/{cid}/{supplier['id']}")
        check("collective summary reconciles to supplier ledger", a["netBalance"] == ledger["netBalance"])
        if args.keep:
            fixture = {"companyId": cid, "billId": bid, "challanId": c["id"], "supplierId": supplier["id"], "linkedBillId": linked["id"]}
            print("FIXTURE " + json.dumps(fixture))
    finally:
        if not args.keep:
            for u in users: call("DELETE", f"/api/users/{u['id']}", expected=(200, 204))
            if other_company: call("DELETE", f"/api/companies/{other_company['id']}", expected=(200, 204))
            call("DELETE", f"/api/companies/{cid}", expected=(200, 204))
            if args.db:
                remaining = sql(args.db, f"SET NOCOUNT ON; SELECT COUNT(*) FROM PrivateCostChanges WHERE CompanyId={cid}")
                check("company deletion clears its private cost history and source links", int(remaining.strip()) == 0)
    print(f"{count}/{count} commercial cost checks passed")


if __name__ == "__main__":
    main()
