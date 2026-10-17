"""Local-only Trader report and workbook regression; creates temporary records."""
import argparse, io, time, json
from datetime import date
from urllib.parse import urlencode, urlparse
from urllib.request import Request,urlopen
from urllib.error import HTTPError
import openpyxl, pyodbc
import test_basic_flows as flows
p=argparse.ArgumentParser(description=__doc__);p.add_argument('--base',required=True);p.add_argument('--database',required=True);a=p.parse_args()
assert urlparse(a.base).hostname in ('localhost','127.0.0.1')
assert a.database.endswith('_Test') and all(c.isalnum() or c=='_' for c in a.database)
db=pyodbc.connect('DRIVER={ODBC Driver 18 for SQL Server};SERVER=localhost\\MSSQLSERVER02;DATABASE='+a.database+';Trusted_Connection=yes;TrustServerCertificate=yes;',autocommit=True)
token,company,client=flows.setup(a.base,'admin','admin123');cid=company['id'];users=[];roles=[];other=None;checks=0

def call(method,path,body=None,auth=None,expected=(200,201,204)):
 if method=='GET' and '/excel' in path and 200 in expected:
  try:
   with urlopen(Request(a.base+path,headers={'Authorization':'Bearer '+(auth or token)})) as response:s,d=response.status,response.read()
  except HTTPError as e:s,d=e.code,e.read()
 else:s,d=flows.http(method,path,a.base,token=auth or token,body=body)
 assert s in expected,(s,path,d);return d

def check(label,ok):
 global checks
 assert ok,label
 checks+=1;print('PASS '+label,flush=True)

def workbook(path):
 req=Request(a.base+path,headers={'Authorization':'Bearer '+token})
 with urlopen(req) as r:return openpyxl.load_workbook(io.BytesIO(r.read()),data_only=False)

def report(params=None,co=None,auth=None):
 return call('GET',f'/api/reports/company/{co or cid}/invoice-sales-detail?'+urlencode(params or {}),auth=auth)

def standalone(buyer,lines,when=None,**extra):
 return call('POST','/api/invoices/standalone',dict(companyId=cid,clientId=buyer,date=when or date.today().isoformat(),gstRate=18,items=lines,**extra))
try:
 other=call('POST','/api/companies',{'name':'[TEMP] Other report company '+str(time.time_ns()),'startingInvoiceNumber':1})
 buyer2=call('POST','/api/clients',{'companyId':cid,'name':'[TEMP] Second report buyer'})
 kind=call('POST','/api/itemtypes',{'companyId':cid,'name':'[TEMP] Hardware category','uom':'Pcs'})
 lines=[{'description':'2mm sheet','quantity':2,'uom':'Pcs','unitPrice':100,'itemTypeId':kind['id']},{'description':'3mm sheet','quantity':3,'uom':'Pcs','unitPrice':50,'itemTypeId':kind['id']}]
 bill=standalone(client['id'],lines,manualChallanNumber='MANUAL-042',freightCharges=25,furtherTaxRate=4)
 otherbill=standalone(buyer2['id'],[{'description':'Other buyer valve','quantity':1,'uom':'Pcs','unitPrice':75}])
 old=standalone(client['id'],[{'description':'Historic goods','quantity':1,'uom':'Pcs','unitPrice':40}],when='2024-01-12')
 # Scope is fixed to synthetic ids created by this test. No provider is called.
 db.execute("UPDATE Invoices SET FbrStatus='Submitted',FbrSubmittedAt=SYSUTCDATETIME(),FbrIRN='LOCAL-TEST-IRN' WHERE Id=? AND CompanyId=?",otherbill['id'],cid)
 r=report();ours=[x for x in r['rows'] if x['invoiceId']==bill['id']]
 check('default period selects this month, not historic bills',len(r['rows'])==3 and all(x['invoiceId']!=old['id'] for x in r['rows']))
 check('original bill descriptions and quantities remain separate',[(x['description'],x['quantity']) for x in ours]==[('2mm sheet',2),('3mm sheet',3)])
 check('manual challan number appears without a system challan',all(x['deliveryChallanNumbers']=='MANUAL-042' for x in ours) and db.execute('SELECT COUNT(*) FROM DeliveryChallans WHERE InvoiceId=?',bill['id']).fetchone()[0]==0)
 check('bill GST and freight-inclusive total reconcile',sum(x['salesTax'] for x in ours)==bill['gstAmount'] and sum(x['total'] for x in ours)==bill['grandTotal']+25)
 check('report counts agree with FBR filters',r['invoiceCount']==2 and r['submittedCount']==1 and r['notSubmittedCount']==1)
 for status,expected in [('submitted',otherbill['id']),('notSubmitted',bill['id'])]:
  filtered=report({'fbrStatus':status});check(status+' filter returns correct complete bills',{x['invoiceId'] for x in filtered['rows']}=={expected})
 filtered=report({'search':'2mm'});check('search retains every line of the matching bill',len(filtered['rows'])==2 and filtered['invoiceCount']==1)
 check('buyer picker retains all period buyers after filtering',len(filtered['buyers'])==2)
 check('client and FBR filters combine',report({'clientId':buyer2['id'],'fbrStatus':'notSubmitted'})['rows']==[])
 check('custom period includes both boundary days',len(report({'period':'custom','from':'2024-01-12','to':'2024-01-12'})['rows'])==1)
 check('legacy month and year still work',len(report({'year':2024,'month':1})['rows'])==1)
 for query in [{'period':'allPeriods'},{'period':'invalid'},{'period':'custom','from':'2026-01-02','to':'2026-01-01'},{'period':'custom','from':'2026-01-01'},{'year':2026,'month':13},{'fbrStatus':'unknown'},{'search':'a'*101},{'period':'custom','from':'9999-12-31','to':'9999-12-31'}]:
  call('GET',f'/api/reports/company/{cid}/invoice-sales-detail?'+urlencode(query),expected=(400,));call('GET',f'/api/reports/company/{cid}/invoice-sales-detail/excel?'+urlencode(query),expected=(400,));check('invalid filter refused for JSON and Excel: '+str(query)[:55],True)
 ws=workbook(f'/api/reports/company/{cid}/invoice-sales-detail/excel?search=2mm').active
 headers=['S. No','Date','Month','DC','DC No','DC #','Inv Series','Inv No','Inv #','Party Name','Address','Ntn','Hs Code','Description','Unit','Qty','Rate','Excl','Tax Rate','G. S. T','Incl','236-G / 236-H Tax','Further Tax','Total','FBR Status','FBR Invoice No','Bill Status']
 check('sales workbook retains Importer column order and labels',[c.value for c in ws[1]]==headers)
 check('workbook filters exactly match displayed bill lines',ws.max_row==4 and [ws.cell(i,14).value for i in (2,3)]==['2mm sheet','3mm sheet'])
 check('sales workbook numeric totals equal report totals',ws.cell(4,18).value==filtered['excludingTax'] and ws.cell(4,24).value==filtered['total'] and ws.cell(2,19).number_format=='0.00%')
 tax=call('GET',f'/api/reports/company/{cid}/tax-sheet?year={date.today().year}');taxrows=[x for x in tax['rows'] if x['invoiceId']==bill['id']]
 check('Tax Sheet screen still groups the existing category',len(taxrows)==1 and taxrows[0]['quantity']==5 and taxrows[0]['itemTypeName']==kind['name'])
 tw=workbook(f'/api/reports/company/{cid}/tax-sheet/excel?year={date.today().year}').active
 th=[tw.cell(5,c).value for c in range(1,12)]
 check('Tax Sheet workbook has category, original description and blank Hscode',th==['NTN Number','Party Name','Inv Number','Inv Date','Item Total QTY','Category for Hscode','Description','Hscode','Excluding Amount','Sales Tax','Total'])
 data=list(tw.iter_rows(min_row=6,max_row=tw.max_row-1,values_only=True));br=[x for x in data if str(x[2])==str(bill['invoiceNumber'])]
 check('Tax Sheet exports each original bill line and its category',len(br)==2 and [x[6] for x in br]==['2mm sheet','3mm sheet'] and all(x[5]==kind['name'] for x in br))
 check('consultant Hscode cells are entirely blank',all(x[7] is None for x in data))
 check('Tax Sheet numeric totals reconcile with grouped screen',sum(x[8] for x in br)==taxrows[0]['excludingAmount'] and sum(x[9] for x in br)==taxrows[0]['salesTax'])
 # Verify cancellation remains visible and included, as in Importer.
 db.execute('UPDATE Invoices SET IsCancelled=1 WHERE Id=? AND CompanyId=?',otherbill['id'],cid)
 cancelled=report();check('cancelled bills remain labelled and included',cancelled['cancelledCount']==1 and cancelled['invoiceCount']==2 and cancelled['total']==r['total'])
 db.execute('UPDATE Invoices SET IsCancelled=0 WHERE Id=? AND CompanyId=?',otherbill['id'],cid)
 # Save a current unclassified filing plan: export must still use original descriptions.
 filing_type=call('POST','/api/itemtypes',{'companyId':cid,'name':'[TEMP] Filing category','uom':'Pcs','hsCode':'9403.6000','fbrUOMId':13,'saleType':'Goods at standard rate (default)'})
 ed=call('GET',f"/api/invoices/{bill['id']}/tax-filing")
 call('PUT',f"/api/invoices/{bill['id']}/tax-filing",{'version':ed['version'],'rows':[{'sourceLineIds':[i['id'] for i in bill['items']],'itemTypeId':filing_type['id'],'description':'Consultant combined category','quantity':5,'valueExclTax':350}]})
 # A legacy partial classification fixture keeps its actual source version.
 plan=json.loads(db.execute('SELECT TaxFilingRowsJson FROM Invoices WHERE Id=? AND CompanyId=?',bill['id'],cid).fetchone()[0]);plan['Rows'][0]['HSCode']=None
 db.execute('UPDATE Invoices SET TaxFilingRowsJson=? WHERE Id=? AND CompanyId=?',json.dumps(plan),bill['id'],cid)
 tw=workbook(f'/api/reports/company/{cid}/tax-sheet/excel?year={date.today().year}').active
 check('saved filing groups do not replace original export descriptions',{tw.cell(i,7).value for i in range(6,tw.max_row)}=={'2mm sheet','3mm sheet'})
 # Spreadsheet text beginning with '=' must never become a formula.
 malicious=standalone(client['id'],[{'description':'=1+2','quantity':1,'uom':'Pcs','unitPrice':10}])
 for suffix,col,row in [('invoice-sales-detail/excel',14,2),('tax-sheet/excel',7,6)]:
  param='search=%3D1%2B2' if col==14 else 'year='+str(date.today().year)
  book=workbook(f'/api/reports/company/{cid}/{suffix}?{param}').active
  cells=[book.cell(i,col) for i in range(row,book.max_row)]
  check('formula-looking description is text in '+suffix,any(c.value and '1+2' in str(c.value) for c in cells) and all(c.data_type!='f' for c in cells))
 for label,keys in [('view',['reports.invoicedetail.view','reports.taxsheet.view']),('export',['reports.invoicedetail.export','reports.taxsheet.export']),('both',['reports.invoicedetail.view','reports.invoicedetail.export','reports.taxsheet.view','reports.taxsheet.export']),('none',[])]:
  name='report_'+label+'_'+str(time.time_ns());u=call('POST','/api/users',{'username':name,'fullName':'Local report tester','password':'LocalTest123!','role':'User'});users.append(u['id'])
  role=call('POST','/api/roles',{'name':name,'permissionKeys':keys,'tenantAdminUserId':u['id']});roles.append(role['id']);call('PUT',f"/api/users/{u['id']}/roles",{'roleIds':[role['id']]});call('PUT',f"/api/usercompanies/user/{u['id']}",{'companyIds':[cid]})
  auth=call('POST','/api/auth/login',{'username':name,'password':'LocalTest123!'})['token']
  for endpoint,permission in [('invoice-sales-detail','reports.invoicedetail.view'),('invoice-sales-detail/excel','reports.invoicedetail.export'),('tax-sheet','reports.taxsheet.view'),('tax-sheet/excel','reports.taxsheet.export')]:
   path=f'/api/reports/company/{cid}/{endpoint}?year={date.today().year}&month={date.today().month}'
   call('GET',path,auth=auth,expected=(200,) if permission in keys else (403,))
   call('GET',f"/api/reports/company/{other['id']}/{endpoint}",auth=auth,expected=(403,))
  check(label+' role: permission and cross-company refusals on all reports',True)
  call('PUT',f"/api/usercompanies/user/{u['id']}",{'companyIds':[]})
  for endpoint in ['invoice-sales-detail','invoice-sales-detail/excel','tax-sheet','tax-sheet/excel']:call('GET',f'/api/reports/company/{cid}/{endpoint}',auth=auth,expected=(403,))
  check(label+' role: revocation immediately blocks all report reads',True)
 check('other company report cannot include fixture bills',report(co=other['id'])['rows']==[])
 print(f'all checks passed ({checks} Trader report checks)',flush=True)
finally:
 for uid in users:call('DELETE',f'/api/users/{uid}')
 for rid in roles:call('DELETE',f'/api/roles/{rid}')
 if other:flows.teardown(a.base,token,other,False)
 flows.teardown(a.base,token,company,False)
 db.close()
