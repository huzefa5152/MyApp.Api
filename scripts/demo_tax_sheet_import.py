"""Local demo: fill a real tax sheet with good and bad rows, preview, commit, show what changed.
Creates and deletes its own company. Not part of the gate.
"""
import io, json, sys, uuid
from datetime import date
import openpyxl
sys.argv = [sys.argv[0], "--base", "http://localhost:5134"]
_src = open("scripts/test_tax_sheet_import.py", encoding="utf-8").read()
exec(_src[:_src.index(chr(10) + "try:" + chr(10))])

try:
    co, cl = make_company("Demo", True)
    cid = co["id"]
    for code in ("8204.1100", "8204.1200", "8516.1010"):
        call("POST", f"/api/itemtypes?companyId={cid}", {"companyId": cid, "name": f"Cat {code} {stamp}", "hsCode": code,
             "uom": "Numbers, pieces, units", "fbrUOMId": 69, "saleType": "Goods at Standard Rate (default)"})
    A = bill(co, cl, [("Pipe wrench 14", 1, 5400), ("Allen key 4mm", 6, 135), ("Allen key 5mm", 6, 175)])   # good, combines
    B = bill(co, cl, [("Electric kettle", 1, 6300)])                                                      # good, qty change
    C = bill(co, cl, [("Tarpal 26x40", 12, 16000), ("Rope 10m", 3, 500)])                                  # one line left blank
    D = bill(co, cl, [("Valve 1.5in", 8, 18000)])                                                         # HS with letters
    E = bill(co, cl, [("Measuring tape", 4, 1650)])                                                       # unknown to FBR
    F = bill(co, cl, [("Fuse 10x38", 12, 570)])                                                           # qty typed, no HS
    blob = export(co)
    ws, head, cols = sheet_rows(blob)
    wb = openpyxl.load_workbook(io.BytesIO(blob)); w = wb.worksheets[0]
    hs, qty, desc = cols["Real Hscode"], cols["Filing Qty (blank = bill qty)"], cols["Real Item Description"]
    num = cols["Inv Number"]
    plan = {
        "Pipe wrench 14": ("8204.1100", None), "Allen key 4mm": ("8204.1200", None), "Allen key 5mm": ("8204.1200", None),
        "Electric kettle": ("8516.1010", "2"),
        "Tarpal 26x40": ("8204.1100", None), "Rope 10m": (None, None),
        "Valve 1.5in": ("84.81ABC", None),
        "Measuring tape": ("9017.2000", None),
        "Fuse 10x38": (None, "12"),
    }
    for r in range(head + 1, w.max_row + 1):
        d = (w.cell(r, cols["Description"]).value or "").strip()
        if d in plan:
            code, q = plan[d]
            if code: w.cell(r, hs).value = code
            if q: w.cell(r, qty).value = q
    out = io.BytesIO(); wb.save(out); data = out.getvalue()

    inv_no = {k: v["invoiceNumber"] for k, v in dict(A=A, B=B, C=C, D=D, E=E, F=F).items()}
    print("invoice numbers:", inv_no)
    s, pre = upload(f"/api/reports/company/{cid}/tax-sheet/import/preview", data)
    print("\n=== PREVIEW ===\n", json.dumps(pre, indent=1))
    s, done = upload(f"/api/reports/company/{cid}/tax-sheet/import/commit", data)
    print("\n=== APPLY ===\n", json.dumps(done, indent=1))
    print("\n=== STATE AFTER ===")
    for k, v in dict(A=A, B=B, C=C, D=D, E=E, F=F).items():
        ed = call("GET", f"/api/invoices/{v['id']}/tax-filing")
        rows = [(r["hsCode"], r["quantity"], len(r["sourceLineIds"])) for r in (ed["rows"] or [])] if ed["hasSavedRows"] else []
        print(f"  invoice {v['invoiceNumber']}: saved={ed['hasSavedRows']} rows(hs,qty,lines)={rows}")
finally:
    for c in made:
        http("DELETE", f"/api/companies/{c}", base, token=token)
