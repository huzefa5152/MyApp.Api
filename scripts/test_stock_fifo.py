#!/usr/bin/env python3
"""
FIFO BY GD stock costing, end to end (2026-09-28).
Rules: CLAUDE.md section 5b-17.

One item held on three GDs, imported through the opening-stock sheet:

    GD        GD date     claimed    qty   value    unit
    FIFO-A    2025-01-10  Jun 2026   10    10,000   1,000
    FIFO-B    2024-05-01  (no)       10     5,000     500   <- OLDEST, but unclaimed
    FIFO-C    2025-03-01  Jul 2026   10    20,000   2,000

A sale this month must take claimed GDs first, oldest GD date first (A, then
C), and only then the unclaimed B, even though B is the oldest of all.

    1. A new company starts on FIFO by GD without anyone switching it.
    2. The read-only compare answers before anything is switched.
    3. The switch is one-way: a company holding stock cannot go back.
    4. Bill pricing returns the tiers in consumption order.
    5. A bill of 15 takes A 10 + C 5, and is costed 20,000 (not the average).
    6. The GD panel says what each GD sold and still holds.
    7. The movement names the GDs it took.
    8. The Excel export gives each GD its own Consumed / Balance.
    9. A second bill crosses into the unclaimed GD.
   10. Selling past everything is never refused.
   11. Deleting that bill returns the stock where it came from.
   12. Still one-way after sales and a restatement; an empty company may
       still choose the average.

Usage:
    python scripts/test_stock_fifo.py [--base URL] [--keep]
"""

import argparse
import io
import json
import sys
from datetime import date, datetime, timedelta, timezone

try:
    import requests
    import openpyxl
except ImportError:
    print("requests and openpyxl are required:  pip install requests openpyxl")
    sys.exit(2)

for _stream in (sys.stdout, sys.stderr):
    try:
        _stream.reconfigure(encoding="utf-8", errors="replace")
    except (AttributeError, ValueError):
        pass

PASS, FAIL = "PASS", "FAIL"
results = []


def check(name, ok, detail=""):
    results.append((PASS if ok else FAIL, name, detail))
    print(f"  [{PASS if ok else FAIL}] {name}" + (f" -- {detail}" if detail and not ok else ""))
    return ok


def near(a, b, tol=0.02):
    return abs(float(a or 0) - float(b or 0)) <= tol


def call(method, url, h, **kw):
    kw.setdefault("timeout", 180)
    return requests.request(method, url, headers=h, **kw)


MAPPING = {
    "sheetSelect": {"mode": "byHeaderText", "mustContain": ["GD Number"]},
    "headerRow": 3,
    "firstDataRow": 4,
    "columns": {
        "claimMonth": 1, "lotRef": 2, "lotDate": 3, "hsCodeShort": 4, "hsCodeFull": 5,
        "itemName": 6, "unit": 9, "balanceQty": 18, "balanceValue": 19,
        "balanceTaxRate": 20, "balanceTax": 21,
    },
    "hsCodeStripSuffix": ":-",
}

GDS = [  # (gd, gd date, claim month text, qty, value)
    ("FIFO-A", "10-01-2025", "Jun 2026", 10, 10000),
    ("FIFO-B", "01-05-2024", None, 10, 5000),
    ("FIFO-C", "01-03-2025", "Jul 2026", 10, 20000),
]


def workbook(item_name):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Jul 2026"
    ws.cell(2, 1, "Stock Sheet")
    for col, text in {1: "Claimed Month", 2: "GD Number", 3: "GD Date", 4: "4 Digit Hs Code",
                      5: "8 Digit Hs code", 6: "Items", 9: "Unit", 18: "Bal Qty",
                      19: "Balance Exl", 20: "Rate", 21: "S.Tax"}.items():
        ws.cell(3, col, text)
    for i, (gd, gd_date, claim, qty, value) in enumerate(GDS):
        r = 4 + i
        if claim:
            ws.cell(r, 1, claim)
        ws.cell(r, 2, gd)
        ws.cell(r, 3, gd_date)
        ws.cell(r, 4, "8481")
        ws.cell(r, 5, "8481.1000:-")
        ws.cell(r, 6, item_name)
        ws.cell(r, 9, "Pcs")
        ws.cell(r, 18, qty)
        ws.cell(r, 19, value)
        ws.cell(r, 20, 0.18)
        ws.cell(r, 21, round(value * 0.18, 2))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def onhand(api, h, cid, item_id):
    r = call("GET", f"{api}/stock/company/{cid}/onhand", h)
    return next((x for x in r.json() if x["itemTypeId"] == item_id), None) if r.ok else None


def gd_rows(api, h, cid, item_id):
    r = call("GET", f"{api}/stock/company/{cid}/gd-details", h, params={"itemTypeId": item_id})
    return {x["gdNumber"]: x for x in r.json()} if r.ok else {}


def movements(api, h, cid, item_id):
    r = call("GET", f"{api}/stock/company/{cid}/movements", h,
             params={"itemTypeId": item_id, "pageSize": 100})
    return r.json().get("items", []) if r.ok else []


def pk_today():
    # The server dates stock-side writes (a restatement) on Pakistan's
    # calendar; bills must be dated the same way or, between midnight and
    # 05:00 PKT, a bill lands on the day before the restatement it follows.
    return (datetime.now(timezone.utc) + timedelta(hours=5)).date()


def bill(api, h, cid, client_id, item, qty, price=3000):
    return call("POST", f"{api}/invoices/standalone", h, json={
        "date": pk_today().strftime("%Y-%m-%dT00:00:00Z"),
        "companyId": cid, "clientId": client_id, "gstRate": 18,
        "items": [{"itemTypeId": item["itemTypeId"], "description": item["itemTypeName"],
                   "quantity": qty, "uom": "Pcs", "unitPrice": price}]})


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--base", default="http://localhost:5134")
    ap.add_argument("--username", default="admin")
    ap.add_argument("--password", default="admin123")
    ap.add_argument("--keep", action="store_true")
    args = ap.parse_args()
    api = args.base.rstrip("/") + "/api"

    r = requests.post(f"{api}/auth/login", timeout=60,
                      json={"username": args.username, "password": args.password})
    if not r.ok:
        print(f"FATAL: login failed ({r.status_code})")
        return 2
    h = {"Authorization": f"Bearer {r.json()['token']}"}

    tag = datetime.now().strftime("%m%d%H%M%S")
    r = call("POST", f"{api}/companies", h, json={
        "name": f"_stock_fifo {tag}", "brandName": "FIFO",
        "fullAddress": "1 Test Street", "phone": "021-0000000", "ntn": "1234567-8",
        "startingChallanNumber": 1, "startingInvoiceNumber": 1,
        "startingPurchaseBillNumber": 1, "startingGoodsReceiptNumber": 1,
        "startingSalesQuoteNumber": 1, "startingSalesOrderNumber": 1,
        "fbrEnabled": False, "inventoryTrackingEnabled": True, "enableGl": False,
    })
    if not r.ok:
        print(f"FATAL: create company failed ({r.status_code} {r.text[:200]})")
        return 2
    cid = r.json()["id"]
    print(f"\nCompany id={cid}  base={args.base}")
    call("POST", f"{api}/stock/company/{cid}/flow-version", h, json={"version": 2})
    r = call("POST", f"{api}/clients", h, json={
        "name": f"FIFO Client {tag}", "address": "1 Test Road, Karachi",
        "phone": "021-1234567", "companyId": cid, "registrationType": "Unregistered"})
    client_id = r.json()["id"] if r.ok else None

    try:
        # ── Import the three-GD sheet ─────────────────────────────────────
        print("\n-- 0. Import one item on three GDs --")
        name = f"FIFO BALL VALVE {tag}"
        book = workbook(name)
        r = requests.post(f"{api}/spreadsheet-import/opening-stock/preview", headers=h,
                          params={"companyId": cid}, data={"mappingJson": json.dumps(MAPPING)},
                          files={"file": ("fifo.xlsx", io.BytesIO(book),
                                          "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
                          timeout=300)
        prev = r.json() if r.ok else {}
        if not check("the sheet previews and can commit", r.ok and prev.get("canCommit"),
                     f"http {r.status_code}: {r.text[:300]}"):
            return report()
        body = {
            "companyId": cid, "fileSha256": prev["fileSha256"], "fileName": "fifo.xlsx",
            "fileSizeBytes": prev["fileSizeBytes"], "asOfDate": date(2026, 7, 1).isoformat(),
            "postInventoryValue": True, "enableInventoryTracking": True,
            "rows": [{k: x.get(k) for k in ("itemName", "hsCode", "isHsCodePartial", "unit",
                                              "quantity", "value", "salesTaxRate", "lotRefs",
                                              "itemTypeId", "lots")}
                     for x in prev["rows"]],
        }
        r = call("POST", f"{api}/spreadsheet-import/opening-stock/commit", h, json=body)
        if not check("the import commits", r.ok, f"http {r.status_code}: {r.text[:300]}"):
            return report()
        rows = call("GET", f"{api}/stock/company/{cid}/onhand", h).json()
        item = next((x for x in rows if x["itemTypeName"].upper().startswith("FIFO BALL VALVE")), None)
        if not check("one item holding 30 worth 35,000", item and near(item["onHand"], 30, 1e-4)
                     and near(item["valueExcludingTax"], 35000),
                     f"{item and (item['onHand'], item['valueExcludingTax'])}"):
            return report()
        iid = item["itemTypeId"]
        g = gd_rows(api, h, cid, iid)
        check("each GD line kept its own GD, date and claim month",
              set(g) == {"FIFO-A", "FIFO-B", "FIFO-C"} and g["FIFO-B"]["claimMonth"] is None
              and (g["FIFO-A"]["claimMonth"] or "").startswith("2026-06"), json.dumps(g)[:300])

        # ── 1. FIFO by default ────────────────────────────────────────────
        print("\n-- 1. A new company starts on FIFO by GD --")
        r = call("GET", f"{api}/stock/company/{cid}/costing-method", h)
        check("the method is GdFifo without anyone switching it",
              r.ok and r.json()["method"] == "GdFifo", r.text[:200])
        check("GD lines already carry the FIFO split",
              all(x.get("remainingQuantity") is not None for x in g.values()))

        # ── 2. Compare ────────────────────────────────────────────────────
        print("\n-- 2. The compare is read-only --")
        r = call("GET", f"{api}/stock/company/{cid}/costing-compare", h)
        c = r.json() if r.ok else {}
        check("compare answers", r.ok, r.text[:200])
        check("with nothing sold the two methods agree",
              near(c.get("weightedAverageValue"), c.get("fifoValue")) and c.get("changedItemCount") == 0,
              json.dumps(c)[:300])

        # ── 3. The switch is one-way ─────────────────────────────────────
        print("\n-- 3. A company holding stock cannot go back to the weighted average --")
        r = call("PUT", f"{api}/stock/company/{cid}/costing-method", h, json={"method": "GdFifo"})
        check("asking for FIFO again changes nothing", r.ok and not r.json().get("changed"), r.text[:200])
        r = call("PUT", f"{api}/stock/company/{cid}/costing-method", h, json={"method": "Nonsense"})
        check("an unknown method is refused", r.status_code == 400, str(r.status_code))
        r = call("PUT", f"{api}/stock/company/{cid}/costing-method", h, json={"method": "WeightedAverage"})
        check("back to the weighted average is refused once stock is held", r.status_code == 400, str(r.status_code))
        check("the method still reads GdFifo",
              call("GET", f"{api}/stock/company/{cid}/costing-method", h).json()["method"] == "GdFifo")
        after = onhand(api, h, cid, iid)
        check("no figure moves before anything is sold",
              near(after["onHand"], 30, 1e-4) and near(after["valueExcludingTax"], 35000))

        # ── 4. Pricing tiers ──────────────────────────────────────────────
        print("\n-- 4. Bill pricing returns the GDs in the order a sale drains them --")
        r = call("GET", f"{api}/invoices/company/{cid}/stock-pricing", h, params={"itemTypeIds": str(iid)})
        p = r.json()[0] if r.ok and r.json() else {}
        tiers = p.get("tiers") or []
        check("tiers: claimed A, claimed C, then unclaimed B",
              [t["gdNumber"] for t in tiers] == ["FIFO-A", "FIFO-C", "FIFO-B"],
              json.dumps(tiers)[:300])
        check("tiers carry each GD's own cost",
              [round(float(t["unitCost"])) for t in tiers] == [1000, 2000, 500])
        check("tiers mark the claimed ones", [t["claimed"] for t in tiers] == [True, True, False])
        check("unit cost is the first tier's", near(p.get("unitCost"), 1000))

        # ── 5. First bill ────────────────────────────────────────────────
        print("\n-- 5. A bill of 15 takes A (10) then C (5) --")
        r = bill(api, h, cid, client_id, item, 15)
        if not check("the bill is created", r.ok, f"http {r.status_code}: {r.text[:300]}"):
            return report()
        after = onhand(api, h, cid, iid)
        check("15 left", near(after["onHand"], 15, 1e-4), str(after["onHand"]))
        check("value left 15,000 (C 5 x 2,000 + B 10 x 500), not the average's 17,500",
              near(after["valueExcludingTax"], 15000), str(after["valueExcludingTax"]))
        check("cost of the sale 20,000", near(after["valueOut"], 20000), str(after["valueOut"]))

        # ── 6. GD panel ───────────────────────────────────────────────────
        print("\n-- 6. Each GD says what it sold and still holds --")
        g = gd_rows(api, h, cid, iid)
        check("A sold 10, holds 0",
              near(g["FIFO-A"]["consumedQuantity"], 10, 1e-4) and near(g["FIFO-A"]["remainingQuantity"], 0, 1e-4))
        check("C sold 5, holds 5 worth 10,000",
              near(g["FIFO-C"]["consumedQuantity"], 5, 1e-4) and near(g["FIFO-C"]["remainingValueExcludingTax"], 10000))
        check("B untouched", near(g["FIFO-B"]["consumedQuantity"], 0, 1e-4) and near(g["FIFO-B"]["remainingQuantity"], 10, 1e-4))

        # ── 7. Movement allocations ──────────────────────────────────────
        print("\n-- 7. The movement names the GDs it took --")
        mv = movements(api, h, cid, iid)
        out = next((m for m in mv if m["direction"] == "Out"), None)
        alloc = {a["gdNumber"]: a for a in (out or {}).get("allocations", [])}
        check("the bill's movement took A 10 and C 5",
              set(alloc) == {"FIFO-A", "FIFO-C"} and near(alloc["FIFO-A"]["quantity"], 10, 1e-4)
              and near(alloc["FIFO-C"]["quantity"], 5, 1e-4), json.dumps(out)[:400] if out else "no Out movement")
        check("both marked claimed", all(a["claimed"] for a in alloc.values()))
        check("the movement's value is the GDs' cost", out and near(out["value"], 20000), str(out and out["value"]))

        # ── 8. Excel ──────────────────────────────────────────────────────
        print("\n-- 8. The export gives each GD its own Consumed / Balance --")
        r = call("GET", f"{api}/stock/company/{cid}/onhand/excel", h)
        if check("the export downloads", r.ok, str(r.status_code)):
            ws = openpyxl.load_workbook(io.BytesIO(r.content), data_only=False).worksheets[0]
            gd_lines = {}
            for row in ws.iter_rows(min_row=4):
                d = row[3].value or ""
                if str(d).startswith("↳"):
                    gd_lines[row[1].value] = row
            check("three ↳ rows, one per GD", set(gd_lines) >= {"FIFO-A", "FIFO-B", "FIFO-C"},
                  str(list(gd_lines)))
            if {"FIFO-A", "FIFO-C"} <= set(gd_lines):
                a, cc = gd_lines["FIFO-A"], gd_lines["FIFO-C"]
                check("A: consumed 10, balance 0", near(a[13].value, 10, 1e-4) and near(a[17].value, 0, 1e-4),
                      f"N={a[13].value} R={a[17].value}")
                check("C: consumed 5 / 10,000, balance 5 / 10,000",
                      near(cc[13].value, 5, 1e-4) and near(cc[14].value, 10000)
                      and near(cc[17].value, 5, 1e-4) and near(cc[18].value, 10000),
                      f"N={cc[13].value} O={cc[14].value} R={cc[17].value} S={cc[18].value}")

        # ── 9. Second bill crosses into the unclaimed GD ─────────────────
        print("\n-- 9. A bill of 10 takes C (5) then the unclaimed B (5) --")
        r = bill(api, h, cid, client_id, item, 10)
        check("the second bill is created", r.ok, r.text[:200])
        after = onhand(api, h, cid, iid)
        check("5 left worth 2,500 (B 5 x 500)",
              near(after["onHand"], 5, 1e-4) and near(after["valueExcludingTax"], 2500),
              f"{after['onHand']} / {after['valueExcludingTax']}")

        # ── 10. Past everything ─────────────────────────────────────────
        print("\n-- 10. Selling past every GD is never refused --")
        r = bill(api, h, cid, client_id, item, 8)
        check("the bill is created (no oversell block)", r.ok, f"http {r.status_code}: {r.text[:200]}")
        third_id = r.json().get("id") if r.ok else None
        after = onhand(api, h, cid, iid)
        check("on-hand goes to -3, as it always has", near(after["onHand"], -3, 1e-4), str(after["onHand"]))
        mv = movements(api, h, cid, iid)
        last_out = next((m for m in mv if m["direction"] == "Out"), None)
        labels = [a["label"] for a in (last_out or {}).get("allocations", [])]
        check("the uncovered 3 are named as not covered by a GD",
              "Not covered by a GD" in labels, str(labels))

        # ── 11. Delete returns the stock ────────────────────────────────
        print("\n-- 11. Deleting that bill puts the stock back where it came from --")
        if third_id:
            r = call("DELETE", f"{api}/invoices/{third_id}", h)
            check("the bill deletes", r.ok, f"http {r.status_code}: {r.text[:200]}")
            after = onhand(api, h, cid, iid)
            check("back to 5 worth 2,500",
                  near(after["onHand"], 5, 1e-4) and near(after["valueExcludingTax"], 2500),
                  f"{after['onHand']} / {after['valueExcludingTax']}")
            g = gd_rows(api, h, cid, iid)
            check("B holds 5 again", near(g["FIFO-B"]["remainingQuantity"], 5, 1e-4),
                  str(g["FIFO-B"].get("remainingQuantity")))

        # ── 11b. Restate to a stock sheet ───────────────────────────────
        print("\n-- 11b. Restating to a stock sheet's GD lines --")
        this_month = pk_today().strftime("%Y-%m-01")
        sheet = [
            {"itemTypeId": iid, "gdNumber": "SHEET-R2", "gdDate": "2025-02-01", "claimMonth": None,
             "sourceRow": 5, "description": "unclaimed line", "quantity": 2, "valueExcludingTax": 400, "salesTaxRate": 18},
            {"itemTypeId": iid, "gdNumber": "SHEET-R1", "gdDate": "2025-06-01", "claimMonth": this_month,
             "sourceRow": 4, "description": "claimed line", "quantity": 3, "valueExcludingTax": 3000, "salesTaxRate": 18},
        ]
        bad = [dict(sheet[0], quantity=9), sheet[1]]
        r = call("POST", f"{api}/stock/company/{cid}/fifo-restatement", h,
                 json={"sourceFile": "sheet.xlsx", "commit": True, "lines": bad})
        body = r.json() if r.ok else {}
        check("a sheet whose quantity is not on-hand is refused (nothing written)",
              r.ok and not body.get("committed") and not body.get("canCommit")
              and body["items"][0].get("error"), r.text[:300])
        r = call("POST", f"{api}/stock/company/{cid}/fifo-restatement", h,
                 json={"sourceFile": "sheet.xlsx", "commit": False, "lines": sheet})
        body = r.json() if r.ok else {}
        check("the preview reports FIFO 2,500 -> sheet 3,400 and can commit",
              r.ok and body.get("canCommit") and not body.get("committed")
              and near(body["items"][0]["fifoValue"], 2500) and near(body["sheetValue"], 3400), r.text[:300])
        check("the preview wrote nothing", near(onhand(api, h, cid, iid)["valueExcludingTax"], 2500))
        r = call("POST", f"{api}/stock/company/{cid}/fifo-restatement", h,
                 json={"sourceFile": "sheet.xlsx", "commit": True, "lines": sheet})
        check("the restatement commits", r.ok and r.json().get("committed"), r.text[:300])
        after = onhand(api, h, cid, iid)
        check("value is the sheet's 3,400, quantity still 5",
              near(after["valueExcludingTax"], 3400) and near(after["onHand"], 5, 1e-4),
              f"{after['onHand']} / {after['valueExcludingTax']}")
        g = gd_rows(api, h, cid, iid)
        check("the GD panel now shows the sheet's two lines, not the old lots",
              set(g) == {"SHEET-R1", "SHEET-R2"}, str(list(g)))
        check("each sheet line holds its own balance",
              near(g["SHEET-R1"]["remainingValueExcludingTax"], 3000) and near(g["SHEET-R2"]["remainingQuantity"], 2, 1e-4))
        r = bill(api, h, cid, client_id, item, 4)
        check("a bill after the restatement is created", r.ok, r.text[:200])
        mv = movements(api, h, cid, iid)
        out = next((m for m in mv if m["direction"] == "Out"), None)
        alloc = {a["gdNumber"]: a for a in (out or {}).get("allocations", [])}
        check("it takes the claimed sheet line first (R1 3, then R2 1)",
              near(alloc.get("SHEET-R1", {}).get("quantity"), 3, 1e-4) and near(alloc.get("SHEET-R2", {}).get("quantity"), 1, 1e-4),
              json.dumps(out)[:300] if out else "no Out")
        check("costed 3,000 + 200", out and near(out["value"], 3200), str(out and out["value"]))
        after = onhand(api, h, cid, iid)
        check("1 left worth 200", near(after["onHand"], 1, 1e-4) and near(after["valueExcludingTax"], 200))

        # ── 12. Still one-way after sales and a restatement ─────────────
        print("\n-- 12. Going back is still refused --")
        r = call("PUT", f"{api}/stock/company/{cid}/costing-method", h, json={"method": "WeightedAverage"})
        check("back to the weighted average is refused", r.status_code == 400, str(r.status_code))
        after = onhand(api, h, cid, iid)
        check("the stock is still valued FIFO (1 left worth 200)",
              near(after["onHand"], 1, 1e-4) and near(after["valueExcludingTax"], 200), str(after["valueExcludingTax"]))
        r = call("GET", f"{api}/invoices/company/{cid}/stock-pricing", h, params={"itemTypeIds": str(iid)})
        check("pricing still offers the FIFO tiers", r.ok and r.json()[0].get("tiers"), r.text[:200])

        # ── 12a. Corrections land exactly on a FIFO company ─────────────
        print("\n-- 12a. A 'set' stock correction lands exactly under FIFO --")
        today = pk_today().isoformat()
        r = call("POST", f"{api}/stock/adjust", h, json={
            "companyId": cid, "itemTypeId": iid, "mode": "set",
            "targetQuantity": 3, "targetValueExcludingTax": 900,
            "movementDate": today, "notes": "recount up"})
        check("a quantity + value correction up is accepted", r.ok, r.text[:200])
        after = onhand(api, h, cid, iid)
        check("it lands on 3 worth 900", near(after["onHand"], 3, 1e-4) and near(after["valueExcludingTax"], 900),
              f"{after['onHand']} / {after['valueExcludingTax']}")
        r = call("POST", f"{api}/stock/adjust", h, json={
            "companyId": cid, "itemTypeId": iid, "mode": "set",
            "targetQuantity": 1, "targetValueExcludingTax": 150,
            "movementDate": today, "notes": "recount down"})
        check("a quantity + value correction down is accepted", r.ok, r.text[:200])
        after = onhand(api, h, cid, iid)
        check("it lands on 1 worth 150 (not where the average would put it)",
              near(after["onHand"], 1, 1e-4) and near(after["valueExcludingTax"], 150),
              f"{after['onHand']} / {after['valueExcludingTax']}")

        # ── 12b. An empty company may still choose ───────────────────────
        print("\n-- 12b. A company that has never held stock may still choose the average --")
        r = call("POST", f"{api}/companies", h, json={
            "name": f"_stock_fifo_empty {tag}", "brandName": "FIFOE",
            "fullAddress": "1 Test Street", "phone": "021-0000000", "ntn": "1234567-8",
            "startingChallanNumber": 1, "startingInvoiceNumber": 1,
            "startingPurchaseBillNumber": 1, "startingGoodsReceiptNumber": 1,
            "startingSalesQuoteNumber": 1, "startingSalesOrderNumber": 1,
            "fbrEnabled": False, "inventoryTrackingEnabled": True, "enableGl": False})
        empty_id = r.json().get("id") if r.ok else None
        if check("a second company is created", empty_id is not None, r.text[:200]):
            check("it starts on FIFO too",
                  call("GET", f"{api}/stock/company/{empty_id}/costing-method", h).json()["method"] == "GdFifo")
            r = call("PUT", f"{api}/stock/company/{empty_id}/costing-method", h, json={"method": "WeightedAverage"})
            check("with nothing held it may choose the weighted average", r.ok and r.json().get("changed"), r.text[:200])
            call("DELETE", f"{api}/companies/{empty_id}", h)

        # ── Isolation ────────────────────────────────────────────────────
        # Cross-tenant refusal of the new routes is proven by
        # scripts/test_tenant_isolation.py with a restricted user: the seed
        # admin this suite logs in as is granted every company by design.
        print("\n-- 13. Switching a company that does not exist is refused --")
        rr = call("PUT", f"{api}/stock/company/999999/costing-method", h, json={"method": "GdFifo"})
        check("PUT costing-method on a company that does not exist is refused",
              rr.status_code in (403, 404), str(rr.status_code))
    finally:
        if args.keep:
            print(f"\nKeeping company id={cid} (--keep)")
        else:
            d = call("DELETE", f"{api}/companies/{cid}", h)
            print(f"\nTeardown: delete company returned {d.status_code}")
    return report()


def report():
    print("\n" + "=" * 72)
    p = sum(1 for s, _, _ in results if s == PASS)
    f = sum(1 for s, _, _ in results if s == FAIL)
    for st, name, detail in results:
        if st == FAIL:
            print(f"  FAIL  {name} -- {detail}")
    print(f"{p}/{p + f} checks passed")
    print("=" * 72)
    return 0 if f == 0 else 1


if __name__ == "__main__":
    sys.exit(main())
