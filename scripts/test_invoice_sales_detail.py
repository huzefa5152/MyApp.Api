#!/usr/bin/env python3
"""
Invoice Sales Detail -- the period presets, the three filters, and the Excel
that must keep its shape.

Builds two throwaway companies. Company A gets bills this month (one of them
two lines long), last month (then cancelled), earlier this year and last year,
for two buyers; company B gets one bill. Then:

  1. periods   no period = This Month; every preset and a custom range select
               exactly the bills dated inside them; old year+month links work
  2. refusals  All Periods, an unknown preset, a half or backwards range, a bad
               month, an unknown FBR status, an over-long search
  3. filters   customer, FBR status and a BILL-level search, alone and
               together; the buyer list is the whole period's; counts and
               totals are the listed rows'
  4. Excel     the same rows as the screen for the same query, today's 27
               headers, today's file name for a whole month and the dates for
               any other range; one month asked two ways is the same workbook
  5. access    view without export cannot download; another company is 403

Usage:
  python scripts/test_invoice_sales_detail.py [--base URL] [--db "<LOCAL conn>"]

--db marks one bill Submitted to FBR in the LOCAL database, so the FBR status
filter is tested on both sides; without it those checks are skipped.
Everything the suite creates is deleted at the end.
"""
from __future__ import annotations

import argparse
import calendar
import io
import json
import sys
import urllib.error
import urllib.parse
import urllib.request
from datetime import date, datetime, timedelta, timezone

import openpyxl

PKT = timezone(timedelta(hours=5))

# The workbook's header row exactly as the export writes it. The redesign
# changed the screen only; this pins that the sheet did not move.
EXCEL_HEADERS = [
    "S. No", "Date", "Month", "DC", "DC No", "DC #", "Inv Series",
    "Inv No", "Inv #", "Party Name", "Address", "Ntn", "Hs Code", "Description",
    "Unit", "Qty", "Rate", "Excl", "Tax Rate", "G. S. T", "Incl",
    "236-G / 236-H Tax", "Further Tax", "Total", "FBR Status", "FBR Invoice No", "Bill Status",
]

# Real tariff codes: HS validation is master-first, so an invented code is refused.
HS_VALVE = "8518.2990"
HS_GLASS = "7013.9900"

results: list[tuple[str, str, str]] = []


def http(method, path, base, token=None, body=None, params=None, binary=False):
    url = base.rstrip("/") + path
    if params:
        url += "?" + urllib.parse.urlencode(params)
    data = json.dumps(body).encode() if body is not None else None
    req = urllib.request.Request(url, data=data, method=method)
    req.add_header("Content-Type", "application/json")
    if token:
        req.add_header("Authorization", f"Bearer {token}")
    try:
        with urllib.request.urlopen(req, timeout=120) as r:
            raw = r.read()
            if binary:
                return r.status, raw, dict(r.headers)
            text = raw.decode() if raw else ""
            return r.status, (json.loads(text) if text else None), dict(r.headers)
    except urllib.error.HTTPError as e:
        raw = e.read() if e.fp else b""
        try:
            return e.code, json.loads(raw.decode()), dict(e.headers or {})
        except Exception:
            return e.code, raw[:300], dict(e.headers or {})


def check(name, ok, detail=""):
    results.append(("PASS" if ok else "FAIL", name, detail))
    print(f"  [{'PASS' if ok else 'FAIL'}] {name}" + (f"  ({detail})" if detail and not ok else ""))
    return ok


def skip(name, why):
    results.append(("SKIP", name, why))
    print(f"  [SKIP] {name}  ({why})")


def month_bounds(d: date) -> tuple[date, date]:
    return d.replace(day=1), d.replace(day=calendar.monthrange(d.year, d.month)[1])


def message(data) -> str:
    return str(data.get("message") if isinstance(data, dict) else data)


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--base", default="http://localhost:5134")
    ap.add_argument("--db", default=None, help="LOCAL connection string, to mark one bill Submitted")
    args = ap.parse_args()
    base = args.base

    st, data, _ = http("POST", "/api/auth/login", base, body={"username": "admin", "password": "admin123"})
    if st != 200:
        print(f"FATAL: admin login failed ({st} {data})")
        return 2
    admin = data["token"]

    def api(method, path, body=None, params=None, token=None, binary=False):
        return http(method, path, base, token or admin, body, params, binary)

    stamp = datetime.now().strftime("%m%d%H%M%S")
    today = datetime.now(PKT).date()
    made_companies: list[int] = []
    made_users: list[int] = []
    made_roles: list[int] = []

    def mk_company(label):
        s, co, _ = api("POST", "/api/companies", {
            "name": f"ZZ ISD {label} {stamp}",
            "startingChallanNumber": 1, "startingInvoiceNumber": 1,
            "startingCreditNoteNumber": 1, "startingDebitNoteNumber": 1,
            "startingPurchaseBillNumber": 1, "startingGoodsReceiptNumber": 1,
            "fbrEnabled": False, "inventoryTrackingEnabled": False, "enableGl": False,
        })
        assert s in (200, 201), f"company {label}: {s} {co}"
        made_companies.append(co["id"])
        return co["id"]

    def mk_client(company, name, ntn=None):
        body = {"companyId": company, "name": name, "registrationType": "Unregistered"}
        if ntn:
            body["ntn"] = ntn
        s, c, _ = api("POST", "/api/clients", body)
        assert s in (200, 201), f"client {name}: {s} {c}"
        return c

    def mk_item(company, label, hs):
        s, it, _ = api("POST", "/api/itemtypes", {
            "name": f"ZZ ISD {label} {stamp}", "hsCode": hs, "uom": "Pcs",
            "companyId": company, "isFavorite": True}, params={"companyId": company})
        assert s in (200, 201), f"item {label}: {s} {it}"
        return it["id"]

    def mk_bill(company, client, when: date, lines):
        s, b, _ = api("POST", "/api/invoices/standalone", {
            "companyId": company, "clientId": client["id"], "date": f"{when.isoformat()}T00:00:00Z",
            "gstRate": 18,
            "items": [{"description": d, "quantity": q, "unitPrice": p, "uom": "Pcs", "itemTypeId": it}
                      for d, q, p, it in lines]})
        assert s in (200, 201), f"bill {when}: {s} {b}"
        return b

    try:
        cid, cid_b = mk_company("A"), mk_company("B")
        alpha = mk_client(cid, f"ZZ ISD Alpha {stamp}", ntn="1234567")
        beta = mk_client(cid, f"ZZ ISD Beta {stamp}")
        other = mk_client(cid_b, f"ZZ ISD Other {stamp}")
        valve, glass = mk_item(cid, "Valve", HS_VALVE), mk_item(cid, "Glass", HS_GLASS)
        other_item = mk_item(cid_b, "Other", HS_GLASS)

        this_first, this_last = month_bounds(today)
        last_month = (this_first - timedelta(days=1)).replace(day=15)
        early_year = min(date(today.year, 1, 10), today)
        last_year = date(today.year - 1, 12, 20)

        b1 = mk_bill(cid, alpha, this_first, [(f"ZZ Valve one {stamp}", 10, 500, valve),
                                              (f"ZZ Valve two {stamp}", 4, 200, valve)])
        b2 = mk_bill(cid, beta, today, [(f"ZZ Glass bowl {stamp}", 2, 1500, glass)])
        b3 = mk_bill(cid, alpha, last_month, [(f"ZZ Valve three {stamp}", 1, 1000, valve)])
        b4 = mk_bill(cid, beta, early_year, [(f"ZZ Glass jug {stamp}", 3, 700, glass)])
        b5 = mk_bill(cid, alpha, last_year, [(f"ZZ Valve four {stamp}", 5, 300, valve)])
        mk_bill(cid_b, other, today, [(f"ZZ Other bill {stamp}", 1, 999, other_item)])
        s, r, _ = api("POST", f"/api/invoices/{b3['id']}/cancel", {"reason": "ZZ ISD test"})
        assert s == 200, f"cancel b3: {s} {r}"

        dated = {"b1": (b1, this_first), "b2": (b2, today), "b3": (b3, last_month),
                 "b4": (b4, early_year), "b5": (b5, last_year)}
        key_of = {b["id"]: k for k, (b, _) in dated.items()}

        submitted = False
        if args.db:
            try:
                import pyodbc
                cn = pyodbc.connect(args.db, timeout=30, autocommit=True)
                cn.cursor().execute(
                    "UPDATE Invoices SET FbrStatus = 'Submitted', FbrIRN = ? WHERE Id = ? AND CompanyId = ?",
                    f"ZZISD{stamp}", b2["id"], cid)
                submitted = True
            except Exception as e:
                print(f"  (could not mark a bill Submitted: {str(e)[:160]})")

        def expect(frm: date, to: date) -> list[str]:
            return sorted(k for k, (_, d) in dated.items() if frm <= d <= to)

        def keys(rep) -> list[str]:
            return sorted({key_of.get(r["invoiceId"], "?") for r in (rep or {}).get("rows", [])})

        def report(params, company=None, token=None):
            s, d, _ = api("GET", f"/api/reports/company/{company or cid}/invoice-sales-detail",
                          params=params, token=token)
            return s, d

        def workbook(params):
            s, blob, headers = api("GET", f"/api/reports/company/{cid}/invoice-sales-detail/excel",
                                   params=params, binary=True)
            return s, (openpyxl.load_workbook(io.BytesIO(blob)) if s == 200 else None), headers

        ALL = {"period": "custom", "from": last_year.isoformat(), "to": today.isoformat()}

        # ══ 1. periods ═══════════════════════════════════════════════════════
        print("\n-- 1. periods --")
        s, d = report({})
        check("no period asked is This Month", s == 200 and keys(d) == expect(this_first, this_last),
              f"{s} {keys(d) if s == 200 else message(d)}")
        check("...and the answer names its dates",
              s == 200 and d["from"][:10] == this_first.isoformat() and d["to"][:10] == this_last.isoformat(),
              f"{(d or {}).get('from')} {(d or {}).get('to')}")
        label = f"{this_first.day} {this_first:%b} {this_first.year} – {this_last.day} {this_last:%b} {this_last.year}"
        check("...labelled the way every report labels a period",
              s == 200 and d["periodLabel"] == label, f"{(d or {}).get('periodLabel')} vs {label}")
        check("...with when it was generated", s == 200 and bool(d.get("generatedAt")))

        q_start = date(today.year, ((today.month - 1) // 3) * 3 + 1, 1)
        q_end = month_bounds(date(q_start.year, q_start.month + 2, 1))[1]
        week_start = today - timedelta(days=today.weekday())
        presets = {
            "thisMonth": (this_first, this_last),
            "lastMonth": month_bounds(this_first - timedelta(days=1)),
            "thisQuarter": (q_start, q_end),
            "thisYear": (date(today.year, 1, 1), date(today.year, 12, 31)),
            "lastYear": (date(today.year - 1, 1, 1), date(today.year - 1, 12, 31)),
            "today": (today, today),
            "thisWeek": (week_start, week_start + timedelta(days=6)),
        }
        for name, (frm, to) in presets.items():
            s, d = report({"period": name})
            check(f"{name} lists the bills dated {frm} to {to}",
                  s == 200 and keys(d) == expect(frm, to), f"{s} {keys(d) if s == 200 else message(d)}")

        s, d = report(ALL)
        check("a custom range lists the bills inside it, both ends included",
              s == 200 and keys(d) == ["b1", "b2", "b3", "b4", "b5"], f"{keys(d) if s == 200 else message(d)}")

        s1, legacy = report({"year": today.year, "month": today.month})
        s2, preset = report({"period": "thisMonth"})
        lines = lambda rep: [(r["invoiceId"], r["lineNumber"], r["total"]) for r in rep["rows"]]
        check("an old year + month link still opens that month",
              s1 == 200 and s2 == 200 and lines(legacy) == lines(preset), f"{s1} {s2}")

        day = {"period": "custom", "from": last_month.isoformat(), "to": last_month.isoformat()}
        s, d = report(day)
        rows = d.get("rows", []) if s == 200 else []
        check("a one-day range lists that day's bill", keys(d) == ["b3"], str(keys(d)))
        check("a cancelled bill stays listed and says so",
              len(rows) == 1 and rows[0]["billStatus"] == "Cancelled", str(rows[:1]))
        check("...and still counts in the bills and the totals, as before",
              s == 200 and d["invoiceCount"] == 1 and d["cancelledCount"] == 1
              and abs(d["total"] - float(b3["grandTotal"])) < 0.005,
              f"{d.get('invoiceCount')} {d.get('cancelledCount')} {d.get('total')} vs {b3.get('grandTotal')}")

        # ══ 2. refusals ══════════════════════════════════════════════════════
        print("\n-- 2. refusals --")
        bad = [
            ({"period": "allPeriods"}, "period"),
            ({"period": "sometime"}, "period"),
            ({"period": "custom", "from": today.isoformat()}, "start and end date"),
            ({"period": "custom", "from": today.isoformat(), "to": last_year.isoformat()}, "on or before"),
            ({"year": today.year, "month": 13}, "valid month"),
            ({"period": "thisMonth", "fbrStatus": "maybe"}, "FBR status"),
            ({"period": "thisMonth", "search": "x" * 101}, "100 characters"),
        ]
        for params, words in bad:
            s, d = report(params)
            check(f"refused: {json.dumps(params)[:60]}",
                  s == 400 and words.lower() in message(d).lower(), f"{s} {message(d)[:90]}")
        s, _, _ = workbook({"period": "allPeriods"})
        check("the Excel route refuses the same", s == 400, str(s))

        # ══ 3. filters ═══════════════════════════════════════════════════════
        print("\n-- 3. filters --")
        s, d = report({**ALL, "clientId": alpha["id"]})
        check("a customer lists only that buyer's bills", s == 200 and keys(d) == ["b1", "b3", "b5"], str(keys(d)))
        buyers = d.get("buyers", []) if s == 200 else []
        check("...while the buyer list still offers every buyer in the period",
              sorted(b["name"] for b in buyers) == sorted([alpha["name"], beta["name"]]), str(buyers))
        check("...each with its id and NTN",
              any(b["clientId"] == alpha["id"] and b.get("ntn") == "1234567" for b in buyers), str(buyers))
        s, d = report({**ALL, "clientId": other["id"]})
        check("another company's buyer matches nothing here",
              s == 200 and d["rows"] == [] and d["invoiceCount"] == 0, f"{s} {keys(d)}")
        s, d = report({**ALL, "search": beta["name"]})
        check("search finds a buyer's bills", s == 200 and keys(d) == ["b2", "b4"], str(keys(d)))
        s, d = report({**ALL, "search": f"valve two {stamp}".upper()})
        check("search is case-blind and BILL-level: a second line brings its whole bill",
              s == 200 and keys(d) == ["b1"] and len(d["rows"]) == 2, f"{keys(d)} {len(d.get('rows', []))}")
        s, d = report({**ALL, "search": HS_GLASS})
        check("search finds an HS code", s == 200 and keys(d) == ["b2", "b4"], str(keys(d)))
        s, d = report({**ALL, "search": "1234567"})
        check("search finds an NTN", s == 200 and keys(d) == ["b1", "b3", "b5"], str(keys(d)))
        s, d = report({**ALL, "search": f"no such thing {stamp}"})
        check("a search that matches nothing lists nothing", s == 200 and d["rows"] == [], str(keys(d)))
        s, d = report({**ALL, "clientId": alpha["id"], "search": f"valve three {stamp}"})
        check("filters combine", s == 200 and keys(d) == ["b3"], str(keys(d)))

        s, d = report({**ALL, "fbrStatus": "notSubmitted"})
        want = ["b1", "b3", "b4", "b5"] if submitted else ["b1", "b2", "b3", "b4", "b5"]
        check("FBR status: not submitted", s == 200 and keys(d) == want, str(keys(d)))
        if submitted:
            s, d = report({**ALL, "fbrStatus": "submitted"})
            check("FBR status: submitted", s == 200 and keys(d) == ["b2"], str(keys(d)))
            check("...showing its FBR invoice number",
                  s == 200 and d["rows"] and d["rows"][0]["fbrInvoiceNumber"] == f"ZZISD{stamp}", str(d.get("rows", [])[:1]))
            s, d = report({**ALL, "search": f"zzisd{stamp}"})
            check("search finds an FBR invoice number", s == 200 and keys(d) == ["b2"], str(keys(d)))
        else:
            for n in ("FBR status: submitted", "...showing its FBR invoice number", "search finds an FBR invoice number"):
                skip(n, "needs --db")

        s, d = report(ALL)
        d = d if s == 200 and isinstance(d, dict) else {}
        rows = d.get("rows", [])
        check("counts are the listed bills",
              d.get("invoiceCount") == 5
              and d.get("submittedCount", 0) + d.get("notSubmittedCount", 0) == 5
              and d.get("cancelledCount") == 1 and d.get("submittedCount") == (1 if submitted else 0),
              f"{s} {d.get('invoiceCount')} {d.get('submittedCount')} {d.get('notSubmittedCount')} {d.get('cancelledCount')}")
        for k, name in (("excludingTax", "Excl"), ("salesTax", "G. S. T"), ("advanceTax", "236-G / 236-H"),
                        ("furtherTax", "further tax"), ("total", "grand")):
            check(f"the {name} total is the sum of its rows",
                  bool(rows) and abs(d.get(k, 0) - sum(r[k] for r in rows)) < 0.005,
                  f"{d.get(k)} vs {sum(r[k] for r in rows)}")
        b1_rows = [r for r in rows if r["invoiceId"] == b1["id"]]
        check("a two-line bill is two rows in line order", [r["lineNumber"] for r in b1_rows] == [1, 2], str(b1_rows))
        check("every line's value is its quantity times its rate",
              all(abs(r["excludingTax"] - r["quantity"] * r["rate"]) < 0.005 for r in rows))

        # ══ 4. Excel ═════════════════════════════════════════════════════════
        print("\n-- 4. Excel --")
        for label_, params in [("the whole range", ALL), ("one customer", {**ALL, "clientId": alpha["id"]}),
                               ("a search", {**ALL, "search": beta["name"]}), ("This Month", {"period": "thisMonth"})]:
            _, d = report(params)
            d = d if isinstance(d, dict) and "rows" in d else {"rows": [], "total": 0}
            s, wb, _ = workbook(params)
            ws = wb.worksheets[0] if wb else None
            inv = [ws.cell(r, 8).value for r in range(2, ws.max_row)] if ws else None
            check(f"Excel ({label_}) holds the screen's lines in the screen's order",
                  s == 200 and inv == [r["invoiceNumber"] for r in d["rows"]],
                  f"{s} {inv} vs {[r['invoiceNumber'] for r in d.get('rows', [])]}")
            check(f"Excel ({label_}) totals row is the screen's total",
                  ws is not None and ws.cell(ws.max_row, 14).value == "Total listed bills"
                  and abs((ws.cell(ws.max_row, 24).value or 0) - d["total"]) < 0.005,
                  f"{ws.cell(ws.max_row, 14).value if ws else None} {ws.cell(ws.max_row, 24).value if ws else None}")

        s, wb, headers = workbook({"period": "thisMonth"})
        ws = wb.worksheets[0]
        check("the sheet is still called Sales Detail", ws.title == "Sales Detail", ws.title)
        check("the header row is still the same 27 columns",
              [ws.cell(1, c).value for c in range(1, 28)] == EXCEL_HEADERS and ws.max_column == 27,
              str([ws.cell(1, c).value for c in range(1, ws.max_column + 1)]))
        check("the header row is still frozen and filtered", ws.freeze_panes == "A2" and bool(ws.auto_filter.ref),
              f"{ws.freeze_panes} {ws.auto_filter.ref}")
        cd = headers.get("Content-Disposition", "")
        check("a whole month keeps today's file name", f"Invoice-Sales-Detail-{today:%Y-%m}.xlsx" in cd, cd)
        _, _, headers = workbook(ALL)
        cd = headers.get("Content-Disposition", "")
        check("any other range names its dates",
              f"Invoice-Sales-Detail-{last_year.isoformat()}_to_{today.isoformat()}.xlsx" in cd, cd)
        lm = month_bounds(this_first - timedelta(days=1))
        _, _, headers = workbook({"period": "custom", "from": lm[0].isoformat(), "to": lm[1].isoformat()})
        cd = headers.get("Content-Disposition", "")
        check("a custom range that is exactly a month is named for the month",
              f"Invoice-Sales-Detail-{lm[0]:%Y-%m}.xlsx" in cd, cd)
        _, w1, _ = workbook({"year": today.year, "month": today.month})
        _, w2, _ = workbook({"period": "custom", "from": this_first.isoformat(), "to": this_last.isoformat()})
        cells = lambda w: [[c.value for c in row] for row in w.worksheets[0].iter_rows()]
        check("one month asked as year + month or as a range is the same workbook",
              w1 is not None and w2 is not None and cells(w1) == cells(w2))

        # ══ 5. access ════════════════════════════════════════════════════════
        print("\n-- 5. access --")
        role_name, username = f"ZZ ISD viewer {stamp}", f"zzisd{stamp}"
        s, role, _ = api("POST", "/api/roles", {"name": role_name, "description": "invoice sales detail probe (test)",
                                                 "permissionKeys": ["reports.invoicedetail.view"]})
        assert s in (200, 201), f"role {s} {role}"
        made_roles.append(role["id"])
        s, user, _ = api("POST", "/api/users", {"username": username, "password": "test1234", "fullName": username,
                                                 "role": role_name, "roleIds": [role["id"]], "companyIds": [cid]})
        assert s in (200, 201), f"user {s} {user}"
        made_users.append(user["id"])
        s, lg, _ = http("POST", "/api/auth/login", base, body={"username": username, "password": "test1234"})
        assert s == 200, f"viewer login {s} {lg}"
        viewer = lg["token"]
        s, _ = report({"period": "thisMonth"}, token=viewer)
        check("a viewer opens their company's report", s == 200, str(s))
        s, _, _ = api("GET", f"/api/reports/company/{cid}/invoice-sales-detail/excel",
                      params={"period": "thisMonth"}, token=viewer, binary=True)
        check("...cannot download it without the export permission", s == 403, str(s))
        s, _ = report({"period": "thisMonth"}, company=cid_b, token=viewer)
        check("...and cannot open another company's", s == 403, str(s))
    finally:
        for uid in made_users:
            api("DELETE", f"/api/users/{uid}")
        for rid in made_roles:
            api("DELETE", f"/api/roles/{rid}")
        for c in made_companies:
            api("DELETE", f"/api/companies/{c}")

    passed = sum(1 for r in results if r[0] == "PASS")
    failed = [r for r in results if r[0] == "FAIL"]
    skipped = sum(1 for r in results if r[0] == "SKIP")
    print(f"\n{passed}/{passed + len(failed)} checks passed" + (f" ({skipped} skipped)" if skipped else ""))
    for _, name, detail in failed:
        print(f"  FAILED: {name}  {detail}")
    return 1 if failed else 0


if __name__ == "__main__":
    sys.exit(main())
