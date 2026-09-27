"""
Onboarding import — one workbook brings a new company's customers, items,
suppliers and opening stock in.

What this asserts, against a running backend:
  1. the sample: the chosen sheets only, Start Here + a hidden Lists sheet,
     the schema's headings, a help row, dropdowns
  2. preview: exact counts per sheet and the reason on every refused row
  3. commit: creates exactly the importable rows, through the normal create
     paths (NTN normalised, province stored as its code, STRN kept only for
     Registered, sites joined with ;)
  4. the same file again creates nothing — every row reads as existing
  5. opening stock can name an item from the Items sheet of the same file
  6. the fix list holds only the refused rows and uploads straight back
  7. permissions: the feature key, and each sheet's own create key
  8. whole-file refusals: over 5,000 rows, a non-spreadsheet
It builds a throwaway company, users and roles and deletes them afterwards.

Usage:
    python scripts/test_onboarding_import.py --base http://localhost:5134
"""

import argparse
import io
import json
import re
import sys
import time
import urllib.error
import urllib.request
import uuid

import zipfile

import openpyxl

BASE = "http://localhost:5134"
PW = "Onboard#Test1"
PASSED, FAILED = [], []
COMPANY_NAME = "ZZ Onboarding Test Co"
TEMP_USERS = ["tempOnbItemsOnly", "tempOnbNoFeature"]
TEMP_ROLES = ["[TEMP] Onboarding items only", "[TEMP] Onboarding no feature"]


def check(suite, label, ok, detail=""):
    (PASSED if ok else FAILED).append((suite, label, detail))
    print("  [%s] %-62s %s" % ("PASS" if ok else "FAIL", label, "" if ok else detail))
    return ok


def request(method, path, token=None, body=None, raw=None, content_type=None):
    headers = {}
    data = None
    if raw is not None:
        data = raw
        headers["Content-Type"] = content_type
    elif body is not None:
        data = json.dumps(body).encode()
        headers["Content-Type"] = "application/json"
    if token:
        headers["Authorization"] = "Bearer " + token
    req = urllib.request.Request(BASE + path, data=data, headers=headers, method=method)
    try:
        with urllib.request.urlopen(req) as resp:
            payload = resp.read()
            ctype = resp.headers.get("Content-Type", "")
            if "json" in ctype:
                return resp.status, json.loads(payload.decode())
            return resp.status, payload
    except urllib.error.HTTPError as e:
        payload = e.read()
        try:
            return e.code, json.loads(payload.decode())
        except Exception:
            return e.code, payload


def upload(path, token, xlsx_bytes, sheets=None, filename="data.xlsx"):
    boundary = "----onb" + uuid.uuid4().hex
    buf = io.BytesIO()
    if sheets is not None:
        buf.write(f"--{boundary}\r\nContent-Disposition: form-data; name=\"sheets\"\r\n\r\n{sheets}\r\n".encode())
    buf.write((f"--{boundary}\r\nContent-Disposition: form-data; name=\"file\"; filename=\"{filename}\"\r\n"
               "Content-Type: application/vnd.openxmlformats-officedocument.spreadsheetml.sheet\r\n\r\n").encode())
    buf.write(xlsx_bytes)
    buf.write(f"\r\n--{boundary}--\r\n".encode())
    body = buf.getvalue()
    # Uploads share the app's "import" rate limit (10 a minute per user). A
    # real onboarding needs three to five; this suite needs more, so it waits.
    for _ in range(3):
        status, data = request("POST", path, token=token, raw=body,
                               content_type=f"multipart/form-data; boundary={boundary}")
        if status != 429:
            return status, data
        print("    (rate limited, waiting 61s)")
        time.sleep(61)
    return status, data


def login(username, password):
    s, d = request("POST", "/api/auth/login", body={"username": username, "password": password})
    assert s == 200, f"login {username}: {s} {d}"
    return d["token"]


def headings(ws):
    return [c.value for c in ws[1] if c.value is not None]


def fill(ws, rows):
    """rows: list of dicts keyed by heading (without the ' *' marker)."""
    col_of = {}
    for c in ws[1]:
        if c.value:
            col_of[str(c.value).replace(" *", "").strip()] = c.column
    r = 3
    for row in rows:
        if row is not None:
            for heading, value in row.items():
                ws.cell(row=r, column=col_of[heading], value=value)
        r += 1


def to_bytes(wb):
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def sheet(preview, key):
    return next(s for s in preview["sheets"] if s["key"] == key)


def row_at(s, n):
    return next((r for r in s["rows"] if r["rowNumber"] == n), None)


def issue_text(r):
    return " | ".join(f"{i.get('column')}: {i['message']}" for i in r["issues"]) if r else "(no row)"


def cleanup(admin):
    _, companies = request("GET", "/api/companies", token=admin)
    for c in companies or []:
        if c["name"].startswith(COMPANY_NAME):
            request("DELETE", f"/api/companies/{c['id']}", token=admin)
    _, users = request("GET", "/api/users", token=admin)
    for u in users or []:
        if u["username"] in TEMP_USERS:
            request("DELETE", f"/api/users/{u['id']}", token=admin)
    _, roles = request("GET", "/api/roles", token=admin)
    for r in roles or []:
        if r["name"] in TEMP_ROLES:
            request("DELETE", f"/api/roles/{r['id']}", token=admin)


def main():
    global BASE
    ap = argparse.ArgumentParser()
    ap.add_argument("--base", default=BASE)
    BASE = ap.parse_args().base.rstrip("/")

    admin = login("admin", "admin123")
    cleanup(admin)

    s, co = request("POST", "/api/companies", token=admin, body={
        "name": COMPANY_NAME, "cnic": "1234567890123", "fbrEnvironment": "sandbox",
        "startingChallanNumber": 1, "startingInvoiceNumber": 1,
    })
    assert s in (200, 201), f"create company: {s} {co}"
    cid = co["id"]
    base_path = f"/api/onboarding-import/company/{cid}"
    print(f"  company {cid} created")

    try:
        # ── 1. Sample ──────────────────────────────────────────────────────
        print("\n  Suite 1 — sample workbook")
        s, blob = request("GET", f"{base_path}/sample", token=admin)
        check("1", "sample downloads", s == 200 and isinstance(blob, bytes) and blob[:2] == b"PK", f"{s}")
        wb = openpyxl.load_workbook(io.BytesIO(blob))
        check("1", "sheets: Start Here, Lists, the four data sheets",
              wb.sheetnames == ["Start Here", "Lists", "Items", "Customers", "Suppliers", "Opening Stock"], str(wb.sheetnames))
        check("1", "Lists sheet is hidden", wb["Lists"].sheet_state == "hidden", wb["Lists"].sheet_state)
        check("1", "Customers headings", headings(wb["Customers"]) == [
            "Name *", "Registration Type *", "NTN", "CNIC", "Province *", "Address", "STRN", "Phone",
            "Sites", "Contact Person"], str(headings(wb["Customers"])))
        check("1", "Items headings", headings(wb["Items"]) == ["Item Name *", "HS Code", "Unit", "Sale Type"],
              str(headings(wb["Items"])))
        check("1", "Suppliers has no Sites or Contact Person",
              "Sites" not in headings(wb["Suppliers"]) and "Contact Person" not in headings(wb["Suppliers"]))
        check("1", "Opening Stock headings", headings(wb["Opening Stock"]) == [
            "Item Name *", "HS Code", "Quantity *", "As Of Date *", "Notes"], str(headings(wb["Opening Stock"])))
        check("1", "help row under the headings", "7 digits" in str(wb["Customers"].cell(row=2, column=3).value or ""))
        # A list that points at another sheet is written as an Excel 2010
        # x14 extension, which openpyxl cannot read — so read the sheet XML.
        z = zipfile.ZipFile(io.BytesIO(blob))
        idx = wb.sheetnames.index("Customers") + 1
        xml = z.read(f"xl/worksheets/sheet{idx}.xml").decode()
        dv = dict((m.group(2).split(":")[0].rstrip("0123456789"), m.group(1))
                  for m in re.finditer(r"<xm:f>(Lists![^<]+)</xm:f>\s*</x14:formula1>\s*<xm:sqref>([^<]+)</xm:sqref>", xml))
        check("1", "Registration Type dropdown reads the Lists sheet", dv.get("B", "").startswith("Lists!$B"), str(dv))
        check("1", "Province dropdown reads the Lists sheet", dv.get("E", "").startswith("Lists!$A"), str(dv))
        check("1", "Lists carries the provinces", wb["Lists"]["A2"].value == "Punjab", str(wb["Lists"]["A2"].value))

        s, only = request("GET", f"{base_path}/sample?sheets=customers", token=admin)
        only_wb = openpyxl.load_workbook(io.BytesIO(only))
        check("1", "?sheets=customers gives only Customers",
              only_wb.sheetnames == ["Start Here", "Lists", "Customers"], str(only_wb.sheetnames))

        # ── 2. Preview ─────────────────────────────────────────────────────
        print("\n  Suite 2 — preview")
        fill(wb["Customers"], [
            {"Name": "ZZ Onb Cust A", "Registration Type": "Registered", "NTN": "1234567-8", "Province": "Sindh",
             "Address": "Plot 1, Karachi", "STRN": "3277876123456", "Phone": "021-1111111"},        # 3 import
            {"Name": "ZZ Onb Cust B", "Registration Type": "unregistered", "Province": "punjab", "STRN": "999"},  # 4 warning
            {"Name": "ZZ Onb Cust C", "Registration Type": "Registered", "Province": "Sindh", "Address": "x"},  # 5 NTN missing
            {"Name": "zz onb cust a ", "Registration Type": "Unregistered", "Province": "Sindh", "Address": "x"},  # 6 dup of 3
            {"Name": "ZZ Onb Cust D", "Registration Type": "CNIC", "CNIC": "42101-1234567-1", "Province": "Islamabad",
             "Address": "F-7, Islamabad", "Sites": "Karachi ; Lahore;", "Contact Person": "Ali; Sara"},  # 7 import
        ])
        fill(wb["Items"], [
            {"Item Name": "ZZ Onb Valve", "HS Code": "8481.8090", "Unit": "Pcs"},   # 3 import
            {"Item Name": "ZZ Onb Service"},                                         # 4 warning (no HS)
            {"Item Name": "ZZ Onb Bad", "HS Code": "84-81"},                         # 5 error
            {"Item Name": "ZZ Onb Valve", "HS Code": "84818090"},                    # 6 dup of 3 (normalised)
            {"Item Name": "ZZ Onb Sale", "HS Code": "8481.8090", "Sale Type": "Nonsense"},  # 7 error
        ])
        fill(wb["Suppliers"], [
            {"Name": "ZZ Onb Supp A", "Registration Type": "Unregistered"},          # 3 import
            {"Name": "ZZ Onb Supp B", "Registration Type": "FTN"},                   # 4 NTN missing
        ])
        fill(wb["Opening Stock"], [
            {"Item Name": "ZZ Onb Valve", "Quantity": 10, "As Of Date": "01-07-2026"},  # 3 import (item from file)
            {"Item Name": "ZZ Onb Missing", "Quantity": 5, "As Of Date": "01-07-2026"},  # 4 no item
            {"Item Name": "ZZ Onb Service", "Quantity": 0, "As Of Date": "01-07-2026"},  # 5 quantity 0
            {"Item Name": "ZZ Onb Valve", "Quantity": 3, "As Of Date": "01-07-2026"},    # 6 dup of 3
        ])
        data = to_bytes(wb)

        s, pv = upload(f"{base_path}/preview", admin, data)
        if not check("2", "preview returns 200", s == 200, f"{s} {pv}"):
            return
        exp = {"customers": (2, 1, 0, 2), "items": (1, 1, 0, 3), "suppliers": (1, 0, 0, 1), "openingStock": (1, 0, 0, 3)}
        for key, (imp, warn, ex, err) in exp.items():
            sh = sheet(pv, key)
            got = (sh["toImport"], sh["withWarnings"], sh["existing"], sh["errors"])
            check("2", f"{key}: import/warning/exists/error = {imp}/{warn}/{ex}/{err}", got == (imp, warn, ex, err), str(got))
        check("2", "total to import is 7", pv["totalToImport"] == 7, str(pv["totalToImport"]))

        cu, it, su, op = (sheet(pv, k) for k in ("customers", "items", "suppliers", "openingStock"))
        check("2", "row 5: NTN required for a Registered customer", "required for a Registered customer" in issue_text(row_at(cu, 5)), issue_text(row_at(cu, 5)))
        check("2", "row 6: duplicate names row 3", "same customer as row 3" in issue_text(row_at(cu, 6)), issue_text(row_at(cu, 6)))
        check("2", "row 4: address warning", "Address" in issue_text(row_at(cu, 4)), issue_text(row_at(cu, 4)))
        check("2", "row 4: STRN not stored warning", "will not be stored" in issue_text(row_at(cu, 4)), issue_text(row_at(cu, 4)))
        check("2", "item row 4: no HS code warning", "cannot be filed" in issue_text(row_at(it, 4)), issue_text(row_at(it, 4)))
        check("2", "item row 5: HS code format", "8481.8090" in issue_text(row_at(it, 5)), issue_text(row_at(it, 5)))
        check("2", "item row 6: 84818090 is the same item as row 3", "row 3" in issue_text(row_at(it, 6)), issue_text(row_at(it, 6)))
        check("2", "item row 7: sale type", "not a sale type" in issue_text(row_at(it, 7)), issue_text(row_at(it, 7)))
        check("2", "supplier row 4: NTN for FTN", "required for a FTN supplier" in issue_text(row_at(su, 4)), issue_text(row_at(su, 4)))
        check("2", "opening row 3 names the file's item", row_at(op, 3)["status"] == "import", issue_text(row_at(op, 3)))
        check("2", "opening row 4: no such item", "no item named" in issue_text(row_at(op, 4)), issue_text(row_at(op, 4)))
        check("2", "opening row 5: quantity > 0", "more than 0" in issue_text(row_at(op, 5)), issue_text(row_at(op, 5)))
        check("2", "opening row 6: same item as row 3", "row 3" in issue_text(row_at(op, 6)), issue_text(row_at(op, 6)))
        s, clients_before = request("GET", f"/api/clients/company/{cid}", token=admin)
        check("2", "preview wrote nothing", s == 200 and not any(c["name"].startswith("ZZ Onb") for c in clients_before))

        # ── 6a. Fix list (before commit) ───────────────────────────────────
        print("\n  Suite 6 — fix list")
        s, fix = upload(f"{base_path}/fix-list", admin, data)
        check("6", "fix list downloads", s == 200 and isinstance(fix, bytes) and fix[:2] == b"PK", f"{s}")
        fwb = openpyxl.load_workbook(io.BytesIO(fix))
        counts = {n: sum(1 for r in fwb[n].iter_rows(min_row=3) if any(c.value for c in r)) for n in fwb.sheetnames}
        check("6", "one row per refused row, per sheet",
              counts == {"Items": 3, "Customers": 2, "Suppliers": 1, "Opening Stock": 3}, str(counts))
        err_col = len(headings(fwb["Customers"]))
        check("6", "Error column says why", "NTN" in str(fwb["Customers"].cell(row=3, column=err_col).value), str(fwb["Customers"].cell(row=3, column=err_col).value))
        s, fpv = upload(f"{base_path}/preview", admin, fix)
        # All nine rows come back. The three that were refused only as a
        # repeat of another row stand alone in the fix list, so they now pass.
        rows_back = sum(len(x["rows"]) for x in fpv["sheets"]) if s == 200 else -1
        errs_back = sum(x["errors"] for x in fpv["sheets"]) if s == 200 else -1
        check("6", "fix list uploads straight back with all nine rows", rows_back == 9, f"{s} {rows_back}")
        check("6", "the six rows with a real problem are still refused", errs_back == 6,
              f"{[(x['key'], x['errors']) for x in fpv['sheets']] if s == 200 else fpv}")

        # ── 3. Commit ──────────────────────────────────────────────────────
        print("\n  Suite 3 — commit")
        s, res = upload(f"{base_path}/commit", admin, data)
        if not check("3", "commit returns 200", s == 200, f"{s} {res}"):
            return
        created = {x["key"]: x["created"] for x in res["sheets"]}
        check("3", "created items 2, customers 3, suppliers 1, opening 1",
              created == {"items": 2, "customers": 3, "suppliers": 1, "openingStock": 1}, str(created))
        check("3", "nothing failed at save", all(x["failed"] == 0 for x in res["sheets"]),
              str([(x["key"], x["failedRows"]) for x in res["sheets"] if x["failed"]]))

        _, clients = request("GET", f"/api/clients/company/{cid}", token=admin)
        by = {c["name"]: c for c in clients}
        a, b, d = by.get("ZZ Onb Cust A"), by.get("ZZ Onb Cust B"), by.get("ZZ Onb Cust D")
        check("3", "customer A: NTN stored as 7 digits", a and a["ntn"] == "1234567", str(a and a["ntn"]))
        check("3", "customer A: province stored as its code (Sindh = 8)", a and a["fbrProvinceCode"] == 8, str(a and a["fbrProvinceCode"]))
        check("3", "customer A: STRN kept for Registered", a and a["strn"] == "3277876123456", str(a and a["strn"]))
        check("3", "customer B: registration type in the list's spelling", b and b["registrationType"] == "Unregistered", str(b and b["registrationType"]))
        check("3", "customer B: STRN dropped for Unregistered", b and not b["strn"], str(b and b["strn"]))
        check("3", "customer D: CNIC stored as digits", d and d["cnic"] == "4210112345671", str(d and d["cnic"]))
        check("3", "customer D: sites joined with ;", d and d["site"] == "Karachi;Lahore", str(d and d["site"]))
        check("3", "customer D: contact persons joined with ;", d and d["contactPerson"] == "Ali;Sara", str(d and d["contactPerson"]))
        check("3", "refused customer C not created", "ZZ Onb Cust C" not in by)

        _, items = request("GET", f"/api/itemtypes?companyId={cid}", token=admin)
        ibn = {(i["name"], i.get("hsCode")): i for i in items}
        valve = ibn.get(("ZZ Onb Valve", "8481.8090"))
        service = ibn.get(("ZZ Onb Service", None))
        check("3", "item with HS code created", valve is not None, str(list(ibn)[:6]))
        check("3", "item without HS code created as a draft", service is not None)
        check("3", "blank sale type took the default", valve and valve["saleType"] == "Goods at standard rate (default)", str(valve and valve["saleType"]))
        check("3", "unit kept", valve and valve["uom"] == "Pcs", str(valve and valve["uom"]))

        _, opening = request("GET", f"/api/stock/company/{cid}/opening", token=admin)
        ov = [o for o in opening or [] if valve and o["itemTypeId"] == valve["id"]]
        check("3", "opening balance on the item created from the same file", len(ov) == 1 and float(ov[0]["quantity"]) == 10, str(ov))
        check("3", "as-of date kept", len(ov) == 1 and ov[0]["asOfDate"].startswith("2026-07-01"), str(ov))

        _, sups = request("GET", f"/api/suppliers/company/{cid}", token=admin)
        check("3", "supplier created", any(x["name"] == "ZZ Onb Supp A" for x in sups or []))

        # ── 4. Same file again ─────────────────────────────────────────────
        print("\n  Suite 4 — the same file again")
        s, pv2 = upload(f"{base_path}/preview", admin, data)
        check("4", "nothing left to import", s == 200 and pv2["totalToImport"] == 0, str(s == 200 and pv2["totalToImport"]))
        ex = {x["key"]: x["existing"] for x in pv2["sheets"]}
        # The created rows, plus each in-file repeat of one: its twin exists now.
        check("4", "every created row (and each repeat of one) reads as existing",
              ex == {"items": 3, "customers": 4, "suppliers": 1, "openingStock": 2}, str(ex))
        s, res2 = upload(f"{base_path}/commit", admin, data)
        check("4", "a second commit creates nothing", s == 200 and res2["totalCreated"] == 0, str(s == 200 and res2["totalCreated"]))

        # ── 5. Opening stock against the catalog alone ─────────────────────
        print("\n  Suite 5 — opening stock names an existing item")
        s, os_sample = request("GET", f"{base_path}/sample?sheets=openingStock", token=admin)
        owb = openpyxl.load_workbook(io.BytesIO(os_sample))
        fill(owb["Opening Stock"], [
            {"Item Name": "zz onb service", "Quantity": 4.5, "As Of Date": "2026-07-01"},  # existing draft item
        ])
        s, pv3 = upload(f"{base_path}/preview", admin, to_bytes(owb), sheets="openingStock")
        op3 = sheet(pv3, "openingStock") if s == 200 else None
        check("5", "resolves an existing catalog item case-insensitively", op3 and op3["toImport"] == 1, issue_text(op3 and op3["rows"][0]))
        check("5", "only the asked sheet is previewed", s == 200 and [x["key"] for x in pv3["sheets"]] == ["openingStock"])

        # ── 5b. NTN matching ───────────────────────────────────────────────
        # One business under a second spelling of its name is still one
        # business: a row whose NTN is already on file is skipped, and two rows
        # sharing an NTN in the file are refused, whatever their names say.
        print("\n  Suite 5b — a known NTN is the same business")
        s, ntn_sample = request("GET", f"{base_path}/sample?sheets=customers,suppliers", token=admin)
        nwb = openpyxl.load_workbook(io.BytesIO(ntn_sample))
        fill(nwb["Customers"], [
            {"Name": "ZZ Onb Cust A Pvt Ltd", "Registration Type": "Registered", "NTN": "1234567", "Province": "Sindh",
             "Address": "x"},                                                                     # 3 exists (A's NTN)
            {"Name": "ZZ Onb Cust E", "Registration Type": "Registered", "NTN": "7654321-0", "Province": "Sindh",
             "Address": "x"},                                                                     # 4 import
            {"Name": "ZZ Onb Cust E Traders", "Registration Type": "Registered", "NTN": "7654321", "Province": "Sindh",
             "Address": "x"},                                                                     # 5 same NTN as row 4
            {"Name": "ZZ Onb Cust F", "Registration Type": "Unregistered", "Province": "Sindh", "Address": "x"},  # 6 import
        ])
        fill(nwb["Suppliers"], [
            {"Name": "ZZ Onb Supp N1", "Registration Type": "Registered", "NTN": "2223334"},    # 3 import
            {"Name": "ZZ Onb Supp N2", "Registration Type": "Registered", "NTN": "2223334-5"},  # 4 same NTN as row 3
        ])
        ntn_data = to_bytes(nwb)
        s, pv5 = upload(f"{base_path}/preview", admin, ntn_data, sheets="customers,suppliers")
        if check("5b", "preview returns 200", s == 200, f"{s} {pv5}"):
            cu5, su5 = sheet(pv5, "customers"), sheet(pv5, "suppliers")
            check("5b", "customers: import/exists/error = 2/1/1",
                  (cu5["toImport"] + cu5["withWarnings"], cu5["existing"], cu5["errors"]) == (2, 1, 1),
                  str((cu5["toImport"], cu5["withWarnings"], cu5["existing"], cu5["errors"])))
            r3 = row_at(cu5, 3)
            check("5b", "a known NTN under another name reads as existing", r3["status"] == "exists", r3["status"])
            check("5b", "and names the customer already holding it",
                  "1234567" in issue_text(r3) and "ZZ Onb Cust A" in issue_text(r3), issue_text(r3))
            check("5b", "two rows with one NTN: the second is refused", "same NTN as row 4" in issue_text(row_at(cu5, 5)),
                  issue_text(row_at(cu5, 5)))
            check("5b", "a row with no NTN is not matched on it", row_at(cu5, 6)["status"] in ("import", "warning"),
                  issue_text(row_at(cu5, 6)))
            check("5b", "suppliers: an 8-digit NTN is the same as its 7 digits",
                  "same NTN as row 3" in issue_text(row_at(su5, 4)), issue_text(row_at(su5, 4)))
        s, res5 = upload(f"{base_path}/commit", admin, ntn_data, sheets="customers,suppliers")
        _, clients5 = request("GET", f"/api/clients/company/{cid}", token=admin)
        names5 = {c["name"] for c in clients5 or []}
        check("5b", "commit creates the new ones only",
              "ZZ Onb Cust E" in names5 and "ZZ Onb Cust F" in names5
              and "ZZ Onb Cust A Pvt Ltd" not in names5 and "ZZ Onb Cust E Traders" not in names5, str(sorted(names5)))

        # ── 7. Permissions ─────────────────────────────────────────────────
        print("\n  Suite 7 — permissions")
        role_ids = []
        for name, keys in ((TEMP_ROLES[0], ["onboarding.import.run", "itemtypes.manage.create"]),
                           (TEMP_ROLES[1], ["clients.manage.create", "itemtypes.manage.create"])):
            s, role = request("POST", "/api/roles", token=admin, body={"name": name, "description": "temp", "permissionKeys": keys})
            assert s in (200, 201), f"role {name}: {s} {role}"
            role_ids.append(role["id"])
        tokens = []
        for username, rid in zip(TEMP_USERS, role_ids):
            s, u = request("POST", "/api/users", token=admin, body={"username": username, "fullName": username, "password": PW, "role": "User"})
            assert s in (200, 201), f"user {username}: {s} {u}"
            request("PUT", f"/api/users/{u['id']}/roles", token=admin, body={"roleIds": [rid]})
            request("PUT", f"/api/usercompanies/user/{u['id']}", token=admin, body={"companyIds": [cid]})
            tokens.append(login(username, PW))
        items_only, no_feature = tokens

        s, _ = request("GET", f"{base_path}/sample?sheets=customers", token=items_only)
        check("7", "asking for a sheet you cannot create is 403", s == 403, str(s))
        s, blob = request("GET", f"{base_path}/sample", token=items_only)
        names = openpyxl.load_workbook(io.BytesIO(blob)).sheetnames if s == 200 else []
        check("7", "no sheets named: only the sheets you can create", names == ["Start Here", "Lists", "Items"], f"{s} {names}")
        s, pv4 = upload(f"{base_path}/preview", items_only, data)
        check("7", "preview covers only the permitted sheet", s == 200 and [x["key"] for x in pv4["sheets"]] == ["items"],
              f"{s} {pv4 if s != 200 else [x['key'] for x in pv4['sheets']]}")
        s, _ = upload(f"{base_path}/commit", items_only, data, sheets="customers")
        check("7", "commit of a forbidden sheet is 403", s == 403, str(s))
        s, _ = request("GET", f"{base_path}/sample", token=no_feature)
        check("7", "without onboarding.import.run: 403", s == 403, str(s))

        # ── 8. Whole-file refusals ─────────────────────────────────────────
        print("\n  Suite 8 — whole-file refusals")
        big = openpyxl.Workbook()
        ws = big.active
        ws.title = "Customers"
        ws.append(["Name", "Registration Type", "Province"])
        for i in range(5001):
            ws.append([f"Bulk {i}", "Unregistered", "Sindh"])
        s, d = upload(f"{base_path}/preview", admin, to_bytes(big), sheets="customers")
        check("8", "over 5,000 rows is refused with the count", s == 400 and "5,001" in json.dumps(d), f"{s} {d}")
        s, d = upload(f"{base_path}/preview", admin, b"name,phone\nx,1\n", filename="data.csv")
        check("8", "a CSV is refused", s == 400, f"{s} {d}")
        s, d = upload(f"{base_path}/preview", admin, b"not a workbook at all", filename="data.xlsx")
        check("8", "an unreadable .xlsx is refused politely", s == 400 and "could not be read" in json.dumps(d), f"{s} {d}")
    finally:
        cleanup(admin)
        _, companies = request("GET", "/api/companies", token=admin)
        check("9", "cleanup: test company removed", not any(c["name"].startswith(COMPANY_NAME) for c in companies or []))

    print()
    total = len(PASSED) + len(FAILED)
    print(f"=== {len(PASSED)}/{total} checks passed ===")
    if FAILED:
        for suite, label, detail in FAILED:
            print(f"  FAIL [{suite}] {label} {detail}")
        sys.exit(1)
    print("all checks passed")


if __name__ == "__main__":
    main()
